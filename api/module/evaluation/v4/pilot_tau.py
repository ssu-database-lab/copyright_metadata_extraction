#!/usr/bin/env python
"""v4 파일럿 200 — 설명 임계값(τ) 보정 도구 (docs/v4/pilot200_tau_calibration_20261010.md).

하위 명령 (저장소 루트에서 실행):
  analyze    런 결과(per_set.jsonl · test_predictions/)로 설명·해시태그·키워드 표, 다른 저작물 쌍(귀무) 분포, τ 스윕, 그림을 만든다.
  judge      독립 LLM(기본 deepseek-v4-pro, DashScope)이 시스템 설명과 정답 설명이 같은 내용인지 맹검 A/B 판정 — 잠정 라벨.
             다른 저작물 쌍(대조군)도 섞어 판정기의 오일치율을 잰다.
  workbook   1단계 사람 검수 워크북: 정답 초안 검수만 (시스템 출력은 보여 주지 않는다 — 정답 오염 방지).
             시트 이름·열 이름이 generate_descriptions.py --import-review 와 호환된다.
  match      2단계 워크북: 1단계 결과로 확정한 정답과 시스템 설명을 A/B 무작위 순서로 보여 주고 일치 판정을 받는다 (출처는 별도 json 에만).
  calibrate  1·2단계가 채워지면 최종 정답과의 코사인을 다시 계산하고 τ 를 고른다 (Youden J 최대, 엄격/관대 두 기준).

예:
  python -m api.module.evaluation.v4.pilot_tau analyze   --run dataset/pilot200_v4/run1 --ids dataset/pilot200_v4/pilot_ids.csv --out dataset/pilot200_v4/analysis
  python -m api.module.evaluation.v4.pilot_tau judge     --run dataset/pilot200_v4/run1 --ids dataset/pilot200_v4/pilot_ids.csv --out dataset/pilot200_v4/analysis
  python -m api.module.evaluation.v4.pilot_tau workbook  --run dataset/pilot200_v4/run1 --ids dataset/pilot200_v4/pilot_ids.csv --out dataset/pilot200_v4/review_1_gold.xlsx
  python -m api.module.evaluation.v4.pilot_tau match     --run dataset/pilot200_v4/run1 --ids dataset/pilot200_v4/pilot_ids.csv --stage1 dataset/pilot200_v4/review_1_gold.xlsx --out dataset/pilot200_v4/review_2_match.xlsx
  python -m api.module.evaluation.v4.pilot_tau calibrate --run dataset/pilot200_v4/run1 --ids dataset/pilot200_v4/pilot_ids.csv --stage1 … --stage2 … --out dataset/pilot200_v4/calibration
"""
from __future__ import annotations

import argparse
import hashlib
import json
import logging
import os
import random
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

import numpy as np
import pandas as pd

log = logging.getLogger("v4.pilot_tau")

ROOT = Path(os.path.abspath(__file__)).parents[4]
DRAFTS = ROOT / "dataset" / "gold_descriptions" / "t002_full" / "drafts.jsonl"
FRAMES = ROOT / "dataset" / "gold_descriptions" / "t002_full" / "frames"
DEFAULT_JUDGE = "deepseek-v4-pro"   # DashScope 경유 — 초안(GPT)·시스템(Qwen·Gemma)과 다른 계열
PROVIDERS = {"dashscope": ("DASHSCOPE_BASE_URL", "https://dashscope-intl.aliyuncs.com/compatible-mode/v1", "DASHSCOPE_API_KEY"),
             "openrouter": (None, "https://openrouter.ai/api/v1", "OPENROUTER_API_KEY")}
JUDGE_PRICE = {"anthropic/claude-sonnet-5.5": (2.0e-6, 1.0e-5)}  # USD / token (in, out), OpenRouter 2026-10-10
TAU_GRID = np.round(np.arange(0.40, 0.9001, 0.01), 2)
LABELS = ("일치", "부분일치", "불일치")
MEDIA_KO = {"image": "이미지", "video": "영상"}

# 1단계 열: generate_descriptions.WB_COLS 의 이름을 그대로 쓴다 (import_review 가 이름으로 읽는다).
STAGE1_COLS = ["work_id", "cell", "media", "source", "bucket", "title", "frames_dir", "work_file", "site_gold_kind", "site_설명(원문)", "site_키워드",
               "draft(GPT)", "draft_hashtags", "draft(Kanana)", "자동검증", "status",
               "verdict", "final_text", "final_hashtags", "reason_codes", "verifier", "verified_at", "notes",
               "tags_content", "entities", "visible_text", "uncertain_points", "unusable_reason", "duration_s", "n_frames"]
STAGE2_COLS = ["pair_id", "cell", "media", "title", "frames_dir", "work_file", "설명 A", "설명 B", "일치판정", "판정_메모", "verifier", "verified_at"]


# ---------------------------------------------------------------- 공통 적재
def load_jsonl(p: Path) -> list[dict]:
    out = []
    with open(p, encoding="utf-8") as f:
        for ln in f:
            ln = ln.strip()
            if ln:
                out.append(json.loads(ln))
    return out


def load_run(run: Path) -> tuple[dict, dict, Optional[dict]]:
    rows: dict[str, dict] = {}
    for r in load_jsonl(run / "per_set.jsonl"):
        rows.setdefault(str(r["set_id"]), {})[r["key"]] = r
    preds = {}
    for p in sorted((run / "test_predictions").glob("*.json")):
        d = json.loads(p.read_text(encoding="utf-8"))
        preds[str(d["set_id"])] = d
    m = run / "test_metrics.json"
    metrics = json.loads(m.read_text(encoding="utf-8")) if m.exists() else None
    return rows, preds, metrics


def load_ids(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    return df.set_index("work_id", drop=False)


def load_drafts(path: Path = DRAFTS) -> dict[str, dict]:
    d: dict[str, dict] = {}
    for r in load_jsonl(path):
        d[str(r["work_id"])] = r  # 마지막 기록이 최신
    return d


def desc_table(rows: dict, preds: dict, ids: pd.DataFrame) -> pd.DataFrame:
    out = []
    for sid in ids.index:
        r = rows.get(sid, {}).get("description")
        p = preds.get(sid) or {}
        meta = p.get("_meta") or {}
        det = (r or {}).get("detail") or {}
        out.append({
            "set_id": sid, "cell": ids.loc[sid, "cell"], "sub_source": ids.loc[sid, "sub_source"], "media": ids.loc[sid, "cell"].split("/")[0],
            "title": ids.loc[sid, "title"], "ok": bool(meta.get("ok")) if p else False, "error": meta.get("error") if p else "예측 없음",
            "kind": (r or {}).get("kind"), "state": (r or {}).get("state"), "pred": p.get("description") or "", "gold": (r or {}).get("gold") or "",
            "cosine": det.get("cosine"), "content_ok": det.get("content_ok"), "shared": det.get("shared_tokens"), "matched": (r or {}).get("matched"),
            "hashtags_source": ((meta.get("pipeline_meta") or {}).get("hashtags_source")),
            "backend": ((meta.get("models") or {}).get("model_used")), "elapsed_sec": meta.get("elapsed_sec"), "cost_krw": meta.get("cost_krw"),
        })
    df = pd.DataFrame(out)
    df["valid"] = (df.state == "value") & (df.pred.str.len() > 0) & (df.gold.str.len() > 0)
    return df


def set_table(rows: dict, ids: pd.DataFrame, key: str) -> pd.DataFrame:
    out = []
    for sid in ids.index:
        r = rows.get(sid, {}).get(key)
        if not r:
            continue
        det = r.get("detail") or {}
        out.append({"set_id": sid, "cell": ids.loc[sid, "cell"], "kind": r.get("kind"), "state": r.get("state"), "set_f1": det.get("set_f1"),
                    "tp": r.get("tp"), "fp": r.get("fp"), "fn": r.get("fn"), "n_pred": len(r.get("pred") or []), "n_gold": len(r.get("gold") or [])})
    return pd.DataFrame(out)


# ---------------------------------------------------------------- 임베딩 · 통계
def make_judge(cache: Optional[Path] = None):
    from .embed_judge import EmbedJudge
    return EmbedJudge(cache_path=cache)


def embed_matrix(judge, texts: list[str]) -> np.ndarray:
    m = np.vstack([np.asarray(judge.embed(t), dtype=np.float64) for t in texts])
    n = np.linalg.norm(m, axis=1, keepdims=True)
    n[n == 0] = 1.0
    return m / n


def null_pairs(P: np.ndarray, G: np.ndarray, cells: list[str], pred_texts: list[str], gold_texts: list[str]) -> dict[str, np.ndarray]:
    """다른 저작물 쌍 (i≠j): 시스템 설명 i 와 정답 j 의 코사인 · 내용 토큰 공유 여부 · 같은 셀 여부."""
    from .embed_judge import content_tokens
    C = P @ G.T
    tp = [set(content_tokens(t)) for t in pred_texts]
    tg = [set(content_tokens(t)) for t in gold_texts]
    n = len(cells)
    cos, ok, same = [], [], []
    for i in range(n):
        for j in range(n):
            if i == j:
                continue
            cos.append(C[i, j])
            ok.append(bool(tp[i] & tg[j]))
            same.append(cells[i] == cells[j])
    return {"cos": np.array(cos), "ok": np.array(ok), "same": np.array(same), "diag": np.diag(C).copy()}


def binary_stats(y_true: np.ndarray, y_pred: np.ndarray) -> dict[str, float]:
    y_true = y_true.astype(bool)
    y_pred = y_pred.astype(bool)
    tp = int((y_true & y_pred).sum()); fp = int((~y_true & y_pred).sum())
    fn = int((y_true & ~y_pred).sum()); tn = int((~y_true & ~y_pred).sum())
    n = tp + fp + fn + tn
    tpr = tp / (tp + fn) if tp + fn else float("nan")
    tnr = tn / (tn + fp) if tn + fp else float("nan")
    po = (tp + tn) / n if n else float("nan")
    pe = (((tp + fp) * (tp + fn)) + ((fn + tn) * (fp + tn))) / (n * n) if n else float("nan")
    kappa = (po - pe) / (1 - pe) if n and pe != 1 else float("nan")
    return {"tp": tp, "fp": fp, "fn": fn, "tn": tn, "tpr": tpr, "fpr": 1 - tnr if tnr == tnr else float("nan"),
            "youden_j": (tpr + tnr - 1) if (tpr == tpr and tnr == tnr) else float("nan"),
            "balanced_acc": (tpr + tnr) / 2 if (tpr == tpr and tnr == tnr) else float("nan"), "accuracy": po, "kappa": kappa}


def tau_sweep(cos: np.ndarray, ok: np.ndarray, null: Optional[dict] = None, labels: Optional[dict[str, np.ndarray]] = None) -> pd.DataFrame:
    """τ 격자마다: 같은 저작물 통과율 · 다른 저작물 쌍 오수락률 · (라벨이 있으면) 라벨 대비 성능."""
    out = []
    for t in TAU_GRID:
        dec = (cos >= t) & ok
        row: dict[str, Any] = {"tau": float(t), "pass_rate": float(dec.mean()) if len(dec) else float("nan")}
        if null is not None and len(null["cos"]):
            fa = (null["cos"] >= t) & null["ok"]
            row["false_accept_all"] = float(fa.mean())
            row["false_accept_same_cell"] = float(fa[null["same"]].mean()) if null["same"].any() else float("nan")
        for name, y in (labels or {}).items():
            mask = ~np.isnan(y)
            st = binary_stats(y[mask] > 0.5, dec[mask])
            for k in ("tpr", "fpr", "youden_j", "balanced_acc", "kappa"):
                row[f"{name}_{k}"] = st[k]
        out.append(row)
    return pd.DataFrame(out)


def select_tau(sweep: pd.DataFrame, col: str, tol: float = 0.02) -> dict[str, Any]:
    """col(예: strict_youden_j) 최대인 τ. 최대값이 여러 τ 에 걸치면 그 구간의 가운데. tol 이내 구간도 보고한다."""
    s = sweep[["tau", col]].dropna()
    if s.empty:
        return {"tau": None, "value": None, "plateau": None, "within_tol": None}
    best = float(s[col].max())
    top = s[np.isclose(s[col], best)].tau.to_numpy()
    near = s[s[col] >= best - tol].tau.to_numpy()
    mid = float(np.round(np.median(top) * 100) / 100)
    return {"tau": mid, "value": best, "plateau": [float(top.min()), float(top.max())], "within_tol": [float(near.min()), float(near.max())]}


def ab_order(key: str) -> bool:
    """True 면 A=정답, B=시스템. 키의 해시로 결정 — 재실행해도 같은 배치."""
    return int(hashlib.sha1(key.encode("utf-8")).hexdigest(), 16) % 2 == 0


# ---------------------------------------------------------------- analyze
def cmd_analyze(a) -> None:
    run, out = Path(a.run), Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    rows, preds, metrics = load_run(run)
    ids = load_ids(Path(a.ids))
    dt = desc_table(rows, preds, ids)
    dt.to_csv(out / "per_set_description.csv", index=False, encoding="utf-8-sig")
    v = dt[dt.valid].reset_index(drop=True)
    judge = make_judge(out / "embed_cache.jsonl")
    P = embed_matrix(judge, v.pred.tolist())
    G = embed_matrix(judge, v.gold.tolist())
    null = null_pairs(P, G, v.cell.tolist(), v.pred.tolist(), v.gold.tolist())
    cos = v.cosine.astype(float).to_numpy()
    ok = v.content_ok.astype(bool).to_numpy()
    labels = None
    jl = out / "judge.jsonl"
    if jl.exists():
        lab = {r["set_id"]: r["label"] for r in load_jsonl(jl) if r.get("kind") == "pair" and r.get("label") in LABELS}
        lv = v.set_id.map(lab)
        labels = {"strict": np.array([1.0 if x == "일치" else (0.0 if x in ("부분일치", "불일치") else np.nan) for x in lv]),
                  "lenient": np.array([1.0 if x in ("일치", "부분일치") else (0.0 if x == "불일치" else np.nan) for x in lv])}
        v["judge_label"] = lv
    sw = tau_sweep(cos, ok, null, labels)
    sw.to_csv(out / "tau_sweep.csv", index=False, encoding="utf-8-sig")
    hs = set_table(rows, ids, "hashtags"); ks = set_table(rows, ids, "keywords")
    hs.to_csv(out / "per_set_hashtags.csv", index=False, encoding="utf-8-sig"); ks.to_csv(out / "per_set_keywords.csv", index=False, encoding="utf-8-sig")
    q = lambda x, p: float(np.quantile(x, p)) if len(x) else None
    same = null["cos"][null["same"]]
    res: dict[str, Any] = {
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "run": str(run), "n_sets": int(len(dt)), "n_ok": int(dt.ok.sum()),
        "errors": dt[~dt.ok][["set_id", "cell", "error"]].to_dict("records"),
        "elapsed_sec": {"sum": float(pd.to_numeric(dt.elapsed_sec, errors="coerce").sum()), "median": float(pd.to_numeric(dt.elapsed_sec, errors="coerce").median())},
        "cost_krw": float(pd.to_numeric(dt.cost_krw, errors="coerce").sum()),
        "backends": dt.groupby(["media", "backend"]).size().reset_index().values.tolist(),
        "hashtags_source": dt.hashtags_source.value_counts(dropna=False).to_dict(),
        "description": {
            "n_valid": int(len(v)), "states": dt.state.value_counts(dropna=False).to_dict(), "kinds": v.kind.value_counts(dropna=False).to_dict(),
            "matched_cos": {"min": q(cos, 0), "p10": q(cos, .1), "median": q(cos, .5), "p90": q(cos, .9), "max": q(cos, 1)},
            "content_ok_rate": float(ok.mean()) if len(ok) else None,
            "null_all": {"n": int(len(null["cos"])), "p95": q(null["cos"], .95), "p99": q(null["cos"], .99), "max": q(null["cos"], 1)},
            "null_same_cell": {"n": int(len(same)), "p95": q(same, .95), "p99": q(same, .99), "max": q(same, 1)},
            "per_cell": [{"cell": c, "n": int(len(g)), "median_cos": float(g.cosine.astype(float).median()),
                          "pass_065": float(((g.cosine.astype(float) >= .65) & g.content_ok.astype(bool)).mean()),
                          "null_same_p99": q(null["cos"][null["same"] & np.array([v.cell[i] == c for i in range(len(v)) for j in range(len(v)) if i != j])], .99)}
                         for c, g in v.groupby("cell")],
            "per_kind": [{"kind": k, "n": int(len(g)), "median_cos": float(g.cosine.astype(float).median()),
                          "pass_065": float(((g.cosine.astype(float) >= .65) & g.content_ok.astype(bool)).mean())} for k, g in v.groupby("kind")],
            "pass_rate_at": {f"{t:.2f}": float(sw.loc[np.isclose(sw.tau, t), "pass_rate"].iloc[0]) for t in (0.55, 0.60, 0.65, 0.70, 0.75)},
            "false_accept_same_cell_at": {f"{t:.2f}": float(sw.loc[np.isclose(sw.tau, t), "false_accept_same_cell"].iloc[0]) for t in (0.55, 0.60, 0.65, 0.70, 0.75)},
        },
        "hashtags": {"mean_set_f1": float(hs.set_f1.astype(float).mean()) if len(hs) else None,
                     "per_cell": hs.groupby("cell").set_f1.mean().round(3).to_dict() if len(hs) else {},
                     "per_kind": hs.groupby(hs.kind.fillna("none")).set_f1.mean().round(3).to_dict() if len(hs) else {},
                     "states": hs.state.value_counts().to_dict() if len(hs) else {}},
        "keywords": {"mean_set_f1": float(ks[ks.state == "value"].set_f1.astype(float).mean()) if len(ks) else None,
                     "per_cell": ks[ks.state == "value"].groupby("cell").set_f1.mean().round(3).to_dict() if len(ks) else {},
                     "states": ks.state.value_counts().to_dict() if len(ks) else {}},
        "metrics": {k: metrics.get(k) for k in ("macro_f1", "micro_f1", "passed", "n_sets", "n_predictions")} if metrics else None,
        "per_key_f1": {k: (vv or {}).get("f1") for k, vv in (metrics or {}).get("keys", {}).items()} if metrics else None,
        "per_cell_macro": {c: ((vv or {}).get("macro") or {}).get("f1") for c, vv in (metrics or {}).get("cells", {}).items()} if metrics else None,
    }
    if labels is not None:
        res["judge"] = {"label_counts": v.judge_label.value_counts(dropna=False).to_dict(),
                        "cos_by_label": {k: {"n": int(len(g)), "median": float(g.cosine.astype(float).median()), "min": float(g.cosine.astype(float).min()), "max": float(g.cosine.astype(float).max())}
                                         for k, g in v.groupby("judge_label")},
                        "tau_strict": select_tau(sw, "strict_youden_j"), "tau_lenient": select_tau(sw, "lenient_youden_j"),
                        "at_065": {n: binary_stats(labels[n][~np.isnan(labels[n])] > .5, ((cos >= .65) & ok)[~np.isnan(labels[n])]) for n in labels}}
        ctl = [r for r in load_jsonl(jl) if r.get("kind") == "control" and r.get("label") in LABELS]
        res["judge"]["controls"] = {"n": len(ctl), "labels": pd.Series([r["label"] for r in ctl]).value_counts().to_dict() if ctl else {}}
        v[["set_id", "cell", "kind", "cosine", "content_ok", "judge_label"]].to_csv(out / "per_set_judge.csv", index=False, encoding="utf-8-sig")
    (out / "analysis.json").write_text(json.dumps(res, ensure_ascii=False, indent=1, default=float), encoding="utf-8")
    plot(out, cos, ok, null, v.get("judge_label"), res)
    log.info("분석 완료 → %s", out)


def plot(out: Path, cos: np.ndarray, ok: np.ndarray, null: dict, judge_label: Optional[pd.Series], res: dict) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    bins = np.linspace(0.2, 1.0, 41)
    fig, ax = plt.subplots(figsize=(7.2, 3.4), dpi=160)
    ax.hist(null["cos"][null["same"]], bins=bins, density=True, color="#b8bcc2", label="different works, same cell (system i vs gold j)")
    if judge_label is not None and judge_label.notna().any():
        styles = {"일치": ("#1f5fbf", "-", "judge: match"), "부분일치": ("#d97b29", "--", "judge: partial"), "불일치": ("#202124", ":", "judge: no match")}
        for lab, (col, ls, name) in styles.items():
            m = (judge_label == lab).to_numpy()
            if m.any():
                ax.hist(cos[m], bins=bins, density=True, histtype="step", linewidth=2, linestyle=ls, color=col, label=f"{name} (n={int(m.sum())})")
    else:
        ax.hist(cos, bins=bins, density=True, histtype="step", linewidth=2, color="#1f5fbf", label=f"same work: system vs gold (n={len(cos)})")
    ax.axvline(0.65, color="#202124", linewidth=1)
    ax.text(0.652, ax.get_ylim()[1] * 0.92, "τ = 0.65", fontsize=8)
    ax.set_xlabel("KURE-v1 cosine")
    ax.set_ylabel("density")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", color="#e6e6e6", linewidth=0.6)
    ax.legend(fontsize=7, frameon=False, loc="upper left")
    fig.tight_layout()
    fig.savefig(out / "cosine_distributions.png")
    plt.close(fig)
    pc = res["description"]["per_cell"]
    fig, ax = plt.subplots(figsize=(7.2, 3.0), dpi=160)
    names = [p["cell"] for p in pc]
    ax.barh(names, [p["pass_065"] for p in pc], color="#1f5fbf", height=0.55)
    for i, p in enumerate(pc):
        ax.text(p["pass_065"] + 0.01, i, f"{p['pass_065']:.0%} (median cos {p['median_cos']:.2f})", va="center", fontsize=7)
    ax.set_xlim(0, 1.15)
    ax.axvline(0.85, color="#202124", linewidth=1, linestyle="--")
    ax.set_xlabel("description pass rate at τ = 0.65 (against draft gold)")
    ax.invert_yaxis()
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(out / "pass_rate_by_cell.png")
    plt.close(fig)


# ---------------------------------------------------------------- judge
JUDGE_SYSTEM = """당신은 저작물 설명문 비교 평가자다. 이미지 또는 영상 한 편을 두고 쓴 것일 수 있는 설명 두 개(A, B)가 주어진다.
두 설명이 같은 화면 내용을 묘사하는지만 판단한다. 문체, 길이, 문장 수, 맞춤법, 어느 쪽이 더 잘 썼는지는 판단하지 않는다.
라벨:
- 일치: 주 피사체, 장소·배경, 주요 행위가 같다. 표현 차이나 사소한 세부(색, 개수, 배경 소품 하나 등)의 차이만 있다.
- 부분일치: 주 피사체는 같지만 중요한 요소가 하나 이상 다르거나 빠졌거나 덧붙었다 (예: 행위가 다름, 장소가 다름, 한쪽에만 있는 주요 사물·인물, 한쪽이 너무 일반적이어서 같은 저작물인지 구별되지 않음).
- 불일치: 주 피사체나 장면이 다르거나 서로 모순된다.
JSON 한 줄로만 답한다: {"label": "일치|부분일치|불일치", "reason": "한 문장"}"""


def _env(name: str) -> Optional[str]:
    v = os.environ.get(name)
    if v:
        return v
    try:
        from dotenv import dotenv_values
        v = (dotenv_values(ROOT / ".env") or {}).get(name)
    except Exception:
        v = None
    return v or None


def _parse_label(text: str) -> Optional[dict]:
    m = re.search(r"\{.*\}", text or "", re.S)
    if not m:
        return None
    try:
        d = json.loads(m.group(0))
    except Exception:
        return None
    return d if d.get("label") in LABELS else None


def cmd_judge(a) -> None:
    from openai import OpenAI
    run, out = Path(a.run), Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    url_env, url_default, key_env = PROVIDERS[a.provider]
    key = _env(key_env)
    if not key:
        raise SystemExit(f"{key_env} 없음 (.env)")
    base = (_env(url_env) if url_env else None) or url_default
    client = OpenAI(base_url=base, api_key=key, timeout=180, max_retries=2)
    rows, preds, _ = load_run(run)
    ids = load_ids(Path(a.ids))
    v = desc_table(rows, preds, ids)
    v = v[v.valid].reset_index(drop=True)
    jobs = []
    for _, r in v.iterrows():
        jobs.append({"kind": "pair", "set_id": r.set_id, "cell": r.cell, "media": r.media, "gold": r.gold, "pred": r.pred, "key": f"pair|{r.set_id}"})
    rng = random.Random(f"pilot200-controls|{a.seed}")
    by_cell: dict[str, list[int]] = {}
    for i, c in enumerate(v.cell):
        by_cell.setdefault(c, []).append(i)
    cells = sorted(by_cell)
    for k in range(a.controls):
        c = cells[k % len(cells)]
        if len(by_cell[c]) < 2:
            continue
        i, j = rng.sample(by_cell[c], 2)
        jobs.append({"kind": "control", "set_id": f"{v.set_id[i]}~{v.set_id[j]}", "cell": c, "media": v.media[i], "gold": v.gold[j], "pred": v.pred[i],
                     "key": f"control|{v.set_id[i]}|{v.set_id[j]}"})
    cache_p = out / "judge.jsonl"
    done = {}
    if cache_p.exists():
        for r in load_jsonl(cache_p):
            if r.get("label") in LABELS and r.get("model") == a.model:
                done[(r["key"], r["text_hash"])] = r
    lock = threading.Lock()
    usage = {"in": 0, "out": 0}

    def text_hash(j):
        return hashlib.sha1((j["gold"] + "\x00" + j["pred"]).encode("utf-8")).hexdigest()[:16]

    def work(j):
        th = text_hash(j)
        if (j["key"], th) in done:
            return done[(j["key"], th)]
        gold_first = ab_order(j["key"])
        A, B = (j["gold"], j["pred"]) if gold_first else (j["pred"], j["gold"])
        user = f"저작물 종류: {MEDIA_KO.get(j['media'], j['media'])}\n\n설명 A:\n{A}\n\n설명 B:\n{B}"
        last = None
        for attempt in range(4):
            try:
                resp = client.chat.completions.create(model=a.model, temperature=0, max_tokens=4000,
                                                      messages=[{"role": "system", "content": JUDGE_SYSTEM}, {"role": "user", "content": user}])
                txt = resp.choices[0].message.content or ""
                d = _parse_label(txt)
                u = getattr(resp, "usage", None)
                if u:
                    with lock:
                        usage["in"] += u.prompt_tokens or 0
                        usage["out"] += u.completion_tokens or 0
                if d:
                    rec = {"kind": j["kind"], "set_id": j["set_id"], "cell": j["cell"], "key": j["key"], "text_hash": th, "model": a.model, "provider": a.provider,
                           "a_is": "gold" if gold_first else "system", "label": d["label"], "reason": d.get("reason", ""), "raw": txt[:500],
                           "at": datetime.now(timezone.utc).isoformat(timespec="seconds")}
                    with lock, open(cache_p, "a", encoding="utf-8") as f:
                        f.write(json.dumps(rec, ensure_ascii=False) + "\n")
                    return rec
                last = f"파싱 실패: {txt[:120]}"
            except Exception as e:  # noqa: BLE001
                last = f"{type(e).__name__}: {e}"
                time.sleep(2 + attempt * 3)
        log.warning("판정 실패 %s — %s", j["key"], last)
        return None

    with ThreadPoolExecutor(max_workers=a.workers) as ex:
        res = list(ex.map(work, jobs))
    got = [r for r in res if r]
    pin, pout = JUDGE_PRICE.get(a.model, (0, 0))
    log.info("판정 %d/%d (쌍 %d · 대조군 %d) · 토큰 in %d out %d · 약 USD %.3f", len(got), len(jobs), sum(r["kind"] == "pair" for r in got),
             sum(r["kind"] == "control" for r in got), usage["in"], usage["out"], usage["in"] * pin + usage["out"] * pout)
    (out / "judge_run.json").write_text(json.dumps({"model": a.model, "provider": a.provider, "n_jobs": len(jobs), "n_ok": len(got), "usage": usage,
                                                    "usd_estimate": usage["in"] * pin + usage["out"] * pout, "controls": a.controls,
                                                    "at": datetime.now(timezone.utc).isoformat(timespec="seconds")}, ensure_ascii=False, indent=1), encoding="utf-8")


# ---------------------------------------------------------------- workbooks
def _clean(v: Any) -> Any:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    if v is None:
        return ""
    if isinstance(v, (list, tuple)):
        v = " ".join(str(x) for x in v)
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
        if v.startswith("="):
            v = "'" + v
    return v


def _style(ws, cols: list[str], widths: dict[str, int], dropdowns: dict[str, str], n_rows: int, freeze: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for i, name in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 12)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for name, formula in dropdowns.items():
        col = get_column_letter(cols.index(name) + 1)
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{max(2, n_rows + 1)}")
        for r in range(2, n_rows + 2):
            ws[f"{col}{r}"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws.freeze_panes = freeze


def cmd_workbook(a) -> None:
    from openpyxl import Workbook
    rows, _, _ = load_run(Path(a.run))
    ids = load_ids(Path(a.ids))
    drafts = load_drafts()
    wb = Workbook()
    ws = wb.active
    ws.title = "review"  # import_review 가 이 이름을 읽는다
    ws.append(STAGE1_COLS)
    n = 0
    for sid in ids.index:
        d = drafts.get(sid, {})
        js = d.get("json") or {}
        spec = d.get("input_spec") or {}
        draft = d.get("draft_text", "") or ""
        tags = " ".join(d.get("draft_hashtags") or [])
        prefill = d.get("status") in ("ok", "needs_review") and bool(draft)
        drow = rows.get(sid, {}).get("description") or {}
        vals = {"work_id": sid, "cell": ids.loc[sid, "cell"], "media": d.get("media"), "source": d.get("source"), "bucket": d.get("bucket"), "title": d.get("title") or ids.loc[sid, "title"],
                "frames_dir": str(FRAMES / sid), "work_file": ids.loc[sid, "work_file"], "site_gold_kind": d.get("site_gold_kind") or drow.get("kind"),
                "site_설명(원문)": d.get("site_raw", ""), "site_키워드": " | ".join(d.get("tags_all") or []), "draft(GPT)": draft, "draft_hashtags": tags,
                "draft(Kanana)": d.get("kanana_text", ""), "자동검증": "; ".join((d.get("checks") or {}).get("problems", [])) or "OK", "status": d.get("status", ""),
                "verdict": "", "final_text": draft if prefill else "", "final_hashtags": tags if prefill else "", "reason_codes": "", "verifier": "", "verified_at": "", "notes": "",
                "tags_content": " | ".join(d.get("tags_content") or []), "entities": json.dumps(js.get("entities_from_title_or_tags", []), ensure_ascii=False),
                "visible_text": json.dumps(js.get("visible_text", []), ensure_ascii=False), "uncertain_points": json.dumps(js.get("uncertain_points", []), ensure_ascii=False),
                "unusable_reason": d.get("unusable_reason", "") or "", "duration_s": spec.get("duration_s", ""), "n_frames": spec.get("n_frames", "")}
        ws.append([_clean(vals.get(c, "")) for c in STAGE1_COLS])
        n += 1
    _style(ws, STAGE1_COLS, {"work_id": 11, "cell": 13, "title": 30, "frames_dir": 28, "work_file": 28, "site_설명(원문)": 40, "site_키워드": 26, "draft(GPT)": 56, "draft_hashtags": 32,
                             "draft(Kanana)": 40, "자동검증": 22, "verdict": 11, "final_text": 56, "final_hashtags": 32, "reason_codes": 12, "notes": 26},
           {"verdict": '"ACCEPT,EDIT,REJECT"'}, n, "C2")
    g = wb.create_sheet("안내")
    for line in [
        "파일럿 200 — 1단계: 정답 초안 검수 (docs/v4/pilot200_tau_calibration_20261010.md)",
        "이 단계에서는 시스템 출력을 보지 않는다. 2단계(일치 판정) 워크북은 1단계가 끝난 뒤 따로 만든다 — 시스템 문장이 정답에 섞이지 않게 하기 위해서다.",
        "1 원본을 연다 (work_file: 이미지는 원본 해상도, 영상 60초 이하는 전체, 60초 초과는 5지점 이상 스크럽). frames_dir 은 초안 모델이 본 프레임이다.",
        "2 verdict: ACCEPT = 수정 없음 / EDIT = 직접 고친 문장을 final_text 에 쓴다 / REJECT = 핵심 오류. draft 열은 고치지 않는다.",
        "3 해시태그: final_hashtags 에 확정 목록 (공백 구분, '#' 포함, 5~10개). 화면에 없는 대상·개인명·라이선스/기증 단어는 지운다.",
        "4 사유 코드 (reason_codes, 쉼표): 1 주피사체·장소·동작 불일치 | 2 화면에 없는 것(환각) | 3 제목·태그 밖 고유명사·개인명 | 4 배경사실·출처·추측·평가어 | 5 형식 | 6 영상 전개 오류·음성 언급 | 7 컬렉션 공통 문구 수준",
        "5 verifier 에 이름 머리글자, verified_at 에 날짜를 쓴다.",
        "끝나면: python -m api.module.evaluation.v4.pilot_tau match --run dataset/pilot200_v4/run1 --ids dataset/pilot200_v4/pilot_ids.csv --stage1 <이 파일> --out dataset/pilot200_v4/review_2_match.xlsx",
    ]:
        g.append([line])
    g.column_dimensions["A"].width = 170
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    log.info("1단계 워크북 %d행 → %s", n, a.out)


def read_stage1(path: Path) -> dict[str, dict]:
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["review"]
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr) if h}
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[col["work_id"]] is None:
            continue
        wid = str(row[col["work_id"]]).strip()
        out[wid] = {k: row[i] for k, i in col.items()}
    return out


def final_gold(sid: str, s1: dict, drafts: dict, desc_row: dict) -> tuple[Optional[str], str]:
    """채점용 최종 정답 설명과 그 출처. 사이트 캡션 정답(kind=caption)은 그대로, 생성 정답은 검수 판정에 따른다."""
    kind = (desc_row or {}).get("kind")
    if kind == "caption":
        return (desc_row.get("gold") or None), "caption"
    r = s1.get(sid) or {}
    verdict = str(r.get("verdict") or "").strip().upper()
    if verdict == "ACCEPT":
        return (drafts.get(sid, {}).get("draft_text") or None), "generated_accept"
    if verdict == "EDIT":
        t = str(r.get("final_text") or "").strip()
        return (t or None), "generated_edit"
    if verdict == "REJECT":
        return None, "generated_reject"
    return None, "pending"


def cmd_match(a) -> None:
    from openpyxl import Workbook
    rows, preds, _ = load_run(Path(a.run))
    ids = load_ids(Path(a.ids))
    drafts = load_drafts()
    s1 = read_stage1(Path(a.stage1))
    wb = Workbook()
    ws = wb.active
    ws.title = "match"
    ws.append(STAGE2_COLS)
    mapping, n, skipped = {}, 0, {}
    for sid in ids.index:
        gold, src = final_gold(sid, s1, drafts, rows.get(sid, {}).get("description") or {})
        pred = (preds.get(sid) or {}).get("description") or ""
        if not gold or not pred:
            skipped[src if not gold else "no_prediction"] = skipped.get(src if not gold else "no_prediction", 0) + 1
            continue
        pid = f"P{n + 1:03d}"
        gold_first = ab_order(f"stage2|{sid}")
        A, B = (gold, pred) if gold_first else (pred, gold)
        mapping[pid] = {"set_id": sid, "a_is": "gold" if gold_first else "system", "gold_source": src}
        ws.append([_clean(x) for x in [pid, ids.loc[sid, "cell"], ids.loc[sid, "cell"].split("/")[0], ids.loc[sid, "title"], str(FRAMES / sid), ids.loc[sid, "work_file"], A, B, "", "", "", ""]])
        n += 1
    _style(ws, STAGE2_COLS, {"pair_id": 8, "cell": 13, "title": 28, "frames_dir": 26, "work_file": 26, "설명 A": 60, "설명 B": 60, "일치판정": 11, "판정_메모": 30},
           {"일치판정": '"일치,부분일치,불일치"'}, n, "G2")
    g = wb.create_sheet("안내")
    for line in [
        "파일럿 200 — 2단계: 두 설명이 같은 화면 내용을 묘사하는지 판정 (A·B 중 어느 쪽이 정답인지는 알리지 않는다).",
        "원본(work_file) 또는 frames_dir 을 보고 판단해도 된다. 문체·길이·맞춤법·잘 쓴 정도는 보지 않는다.",
        "일치: 주 피사체, 장소·배경, 주요 행위가 같다. 표현 차이나 사소한 세부 차이만 있다.",
        "부분일치: 주 피사체는 같지만 중요한 요소가 하나 이상 다르거나 빠졌거나 덧붙었다 (행위·장소가 다름, 한쪽에만 있는 주요 사물·인물, 너무 일반적이어서 구별되지 않음).",
        "불일치: 주 피사체나 장면이 다르거나 서로 모순된다.",
        "끝나면: python -m api.module.evaluation.v4.pilot_tau calibrate --run dataset/pilot200_v4/run1 --ids dataset/pilot200_v4/pilot_ids.csv --stage1 <1단계 파일> --stage2 <이 파일> --out dataset/pilot200_v4/calibration",
    ]:
        g.append([line])
    g.column_dimensions["A"].width = 170
    wb.save(a.out)
    mp = Path(a.out).with_suffix(".mapping.json")
    mp.write_text(json.dumps(mapping, ensure_ascii=False, indent=1), encoding="utf-8")
    log.info("2단계 워크북 %d쌍 → %s (제외 %s; A/B 대응표 %s)", n, a.out, skipped, mp.name)


# ---------------------------------------------------------------- calibrate
def cmd_calibrate(a) -> None:
    from openpyxl import load_workbook
    rows, preds, _ = load_run(Path(a.run))
    ids = load_ids(Path(a.ids))
    drafts = load_drafts()
    s1 = read_stage1(Path(a.stage1))
    mapping = json.loads(Path(a.stage2).with_suffix(".mapping.json").read_text(encoding="utf-8"))
    ws = load_workbook(a.stage2, data_only=True)["match"]
    hdr = [c.value for c in ws[1]]
    ci = {h: i for i, h in enumerate(hdr)}
    human = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        pid = row[ci["pair_id"]]
        lab = str(row[ci["일치판정"]] or "").strip()
        if pid in mapping and lab in LABELS:
            human[mapping[pid]["set_id"]] = lab
    recs = []
    for sid in ids.index:
        gold, src = final_gold(sid, s1, drafts, rows.get(sid, {}).get("description") or {})
        pred = (preds.get(sid) or {}).get("description") or ""
        if gold and pred and sid in human:
            recs.append({"set_id": sid, "cell": ids.loc[sid, "cell"], "gold": gold, "pred": pred, "gold_source": src, "human": human[sid]})
    if not recs:
        raise SystemExit("판정된 쌍이 없다 — 2단계 워크북의 일치판정 열을 채운 뒤 다시 실행")
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame(recs)
    judge = make_judge(out / "embed_cache.jsonl")
    dec = [judge.decide(p, g) for p, g in zip(df.pred, df.gold)]
    df["cosine"] = [d["cosine"] for d in dec]
    df["content_ok"] = [d["content_ok"] for d in dec]
    P = embed_matrix(judge, df.pred.tolist())
    G = embed_matrix(judge, df.gold.tolist())
    null = null_pairs(P, G, df.cell.tolist(), df.pred.tolist(), df.gold.tolist())
    labels = {"strict": np.array([1.0 if x == "일치" else 0.0 for x in df.human]), "lenient": np.array([1.0 if x != "불일치" else 0.0 for x in df.human])}
    cos, ok = df.cosine.to_numpy(), df.content_ok.to_numpy().astype(bool)
    sw = tau_sweep(cos, ok, null, labels)
    sw.to_csv(out / "tau_sweep_human.csv", index=False, encoding="utf-8-sig")
    res = {"n_pairs": len(df), "human_labels": df.human.value_counts().to_dict(), "gold_sources": df.gold_source.value_counts().to_dict(),
           "tau_strict": select_tau(sw, "strict_youden_j"), "tau_lenient": select_tau(sw, "lenient_youden_j"),
           "at_065": {n: binary_stats(labels[n] > .5, (cos >= .65) & ok) for n in labels},
           "cos_by_label": {k: {"n": int(len(g)), "median": float(g.cosine.median()), "min": float(g.cosine.min()), "max": float(g.cosine.max())} for k, g in df.groupby("human")},
           "null_same_cell": {"p95": float(np.quantile(null["cos"][null["same"]], .95)), "p99": float(np.quantile(null["cos"][null["same"]], .99))}}
    jl = Path(a.judge) if a.judge else None
    if jl and jl.exists():
        jlab = {r["set_id"]: r["label"] for r in load_jsonl(jl) if r.get("kind") == "pair" and r.get("label") in LABELS}
        both = df[df.set_id.isin(jlab)]
        if len(both):
            jj = both.set_id.map(jlab)
            res["judge_vs_human"] = {"n": int(len(both)), "exact_agreement": float((jj == both.human).mean()),
                                     "strict_kappa": binary_stats((both.human == "일치").to_numpy(), (jj == "일치").to_numpy())["kappa"],
                                     "lenient_kappa": binary_stats((both.human != "불일치").to_numpy(), (jj != "불일치").to_numpy())["kappa"],
                                     "note": "판정기 라벨은 검수 전 초안 정답 기준이다"}
    df.drop(columns=["gold", "pred"]).to_csv(out / "per_pair_human.csv", index=False, encoding="utf-8-sig")
    (out / "calibration.json").write_text(json.dumps(res, ensure_ascii=False, indent=1, default=float), encoding="utf-8")
    log.info("보정 완료: 엄격 τ=%s · 관대 τ=%s → %s", res["tau_strict"]["tau"], res["tau_lenient"]["tau"], out)


def main(argv: Optional[list[str]] = None) -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("analyze", "judge", "workbook", "match", "calibrate"):
        p = sub.add_parser(name)
        p.add_argument("--run", required=True)
        p.add_argument("--ids", required=True)
        p.add_argument("--out", required=True)
        if name == "judge":
            p.add_argument("--model", default=DEFAULT_JUDGE)
            p.add_argument("--provider", choices=sorted(PROVIDERS), default="dashscope")
            p.add_argument("--controls", type=int, default=40)
            p.add_argument("--workers", type=int, default=6)
            p.add_argument("--seed", type=int, default=20261010)
        if name in ("match", "calibrate"):
            p.add_argument("--stage1", required=True)
        if name == "calibrate":
            p.add_argument("--stage2", required=True)
            p.add_argument("--judge", default=None, help="judge.jsonl (선택) — 판정기와 사람의 일치도")
    a = ap.parse_args(argv)
    {"analyze": cmd_analyze, "judge": cmd_judge, "workbook": cmd_workbook, "match": cmd_match, "calibrate": cmd_calibrate}[a.cmd](a)


if __name__ == "__main__":
    main()
