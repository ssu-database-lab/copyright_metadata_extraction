"""v4 평가 — 23키 Gold 빌더 (docs/v4/M1_evaluator_spec.md §2–§3).

frozen id 목록(dataset/test_set_v4/t00X_ids.csv 또는 1열 id 파일) × 통합 인덱스(gt_merged_index.csv)
× 계약서 원장(생성계약서_메타데이터_항목.xlsx '데이터 랜덤 값') × 생성 설명 초안(drafts.jsonl)
→ 세트별 Gold 레코드(gold23_t00X.jsonl). 레코드 형식은 §2:

  {"set_id", "test", "media", "bucket", "source",
   "keys": {key: {"value", "state": value|empty|na, "source", "tier", "kind"}}}

키별 출처 규칙(§3)과 이 모듈이 확정한 해석:
  * T-001(키 1–16)은 **원장 행만** 출처다 — 인덱스의 contract_gt_* 는 구 템플릿 기준이라 쓰지 않는다.
    원장 행이 없으면(계약서 미생성) 16키 모두 na. 새 템플릿 매핑: work_title←제목(순수 제목),
    creator←저작자, copyright_holder←권리자('-'/빈값→na), 권리 6–12←R_1..R_7(v/V/■=허용, □=불허,
    그 외 글리프→na), other_rights→na(기타 권리 항목이 없는 전 계약서), 14/15←이용허락기간_S*/E*,
    16←'시작~종료', work_type←C_1..C_10 중 체크된 칸(폐쇄 집합). **file_name 은 시스템에 실제로 주는 계약서 PDF 의
    basename**(§1 #2): --file-name-from auto(기본) 는 id 목록의 contract_pdf 가 있으면 그 basename, 없으면 '<원장 파일명>.pdf';
    ids / ledger 는 한쪽만 쓴다. 장세영 쪽 명명 요청('<원문인덱스> <제목>.pdf')이 지켜지면 두 값은 같고, 전달 5건
    ('00002_13387752 운수 자판기 (1).pdf')처럼 다를 때는 contract_pdf 가 정답이다.
  * 원장 파일명은 '<id> <제목>' 에 앞의 생성 순번('00002_')·뒤의 다운로드 표시(' (1)')가 붙어도 id 를 읽는다. id 를 읽은
    행이 0건이거나 절반 미만이면 원장 형식이 바뀐 것이므로 조용히 전부 na 로 가지 않고 ValueError 로 멈춘다.
  * T-002(키 17–23)은 인덱스+생성 Gold 가 출처. T-002 레코드의 1–16, T-001 레코드의 17–23은 비채점
    키이므로 na 로 채운다(file_name 만 양쪽 모두 값: T-002 는 work_file 의 basename).
  * 설명: 사이트 설명(tier A)의 kind 를 dataset_builder.generate_descriptions.classify_site_descriptions
    와 같은 규칙으로 판정한다 — none(URL·귀속 문구) / label(같은 media×bucket×source 에 3행 이상) /
    plot(video·expired) / background(video·공공누리) / caption. **caption 만 사이트 Gold** 로 쓰고,
    나머지는 생성 Gold 가 있으면 생성 Gold(kind generated), 없으면 na. caption 과 생성 Gold 가 모두
    있으면 caption 이 우선한다(사이트 트랙).
  * 해시태그: 공공누리(KOGL) 원본의 work_해시태그 → kind site; 그 외는 생성 Gold(kind generated) 또는 na.
  * 키워드: work_키워드_내용태그 가 있으면 그것, 없고 배치 수준(work_키워드_batch_level=True)이면 na,
    아니면 work_키워드; 둘 다 없으면 na. 태그는 NFC·공백 정리 후 중복(정규화 키 기준)만 제거한다 —
    세밀한 정규화(norm_tag)는 평가기가 양쪽에 똑같이 적용한다.
  * 상태 ②(empty)는 "수집했는데 비어 있음"에만 쓴다: 키워드/해시태그의 tier 가 A 인데 목록이 빈 경우
    (현재 인덱스에는 0건 — 구조적으로만 지원). 그 외의 빈값은 모두 ③(na).
  * 해상도: text 미디어는 na(비적용); 값이 없으면 na. file_format/file_size 는 인덱스 값, 없으면 na.
  * file_hash: --hash 일 때만 SHA-256 을 계산한다(dataset/_probe/file_sha256.jsonl 캐시, path+size+mtime
    키). 아니면 na("미계산") 이며 /mnt/e 의 파일을 건드리지 않는다.
  * 생성 Gold: --generated draft → drafts.jsonl 의 status=ok 초안(tier 'A-검수(대기)');
    --generated verified → gold_descriptions.jsonl(검수 완료, tier 'A-검수')을 우선하고 미검수 세트는
    초안으로 보완, 검수 거절은 na; --generated none → 생성 Gold 를 쓰지 않는다.

CLI (§3): python -m api.module.evaluation.v4.gold23 --test T-002 --ids dataset/test_set_v4/t002_ids.csv
          --out dataset/test_set_v4/gold23_t002.jsonl [--generated draft|verified|none] [--hash]
          [--ledger path] [--file-name-from auto|ids|ledger]   → 끝에 셀×키 커버리지 표(value/empty/na 건수)를 출력한다.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import os
import re
import sys
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping, Optional, Sequence

import pandas as pd

from .keys23 import KEY_NAMES, T001_KEYS, T002_KEYS, WORK_TYPES

log = logging.getLogger("v4.gold23")

STATE_VALUE = "value"
STATE_EMPTY = "empty"
STATE_NA = "na"
GENERATED_MODES: tuple[str, ...] = ("draft", "verified", "none")
TESTS: tuple[str, ...] = ("T-001", "T-002")
FILE_NAME_MODES: tuple[str, ...] = ("auto", "ids", "ledger")  # T-001 file_name gold 출처 (§1 #2)

DEFAULT_INDEX = Path("dataset/gt_merged_index.csv")
DEFAULT_LEDGER = Path("dataset/생성계약서_메타데이터_항목.xlsx")
DEFAULT_LEDGER_SHEET = "데이터 랜덤 값"
DEFAULT_DRAFTS = Path("dataset/gold_descriptions/t002_full/drafts.jsonl")
VERIFIED_BASENAME = "gold_descriptions.jsonl"
DEFAULT_HASH_CACHE = Path("dataset/_probe/file_sha256.jsonl")

# ---------------------------------------------------------------- 원장(ledger) 열 정의 — 새 템플릿(장세영 2026-09-22+)
LEDGER_KNOWN_COLUMNS: frozenset[str] = frozenset(
    ["공유형태", "저작물형태", "파일명", "권리자", "이용자", "저작권자명", "제목", "저작자", "C_기타",
     "R_3_1", "R_3_2", "R_3_3", "양도_SY", "양도_SM", "양도_SD", "Ls_1", "Ls_2", "Ls_3", "정기지급액",
     "Po_1", "Po_2", "Po_3", "분할1_MO", "분할2_MO", "분할3_MO", "정기지급_M", "정기지급_B", "정기지급_Y",
     "지급시기_ETC", "권리자_생년월일", "권리자_주소", "이용자_생년월일", "이용자_주소"]
    + [f"C_{i}" for i in range(1, 11)] + [f"R_{i}" for i in range(1, 8)]
    + [f"이용허락기간_{p}{u}" for p in "SE" for u in "YMD"]
)
LEDGER_REQUIRED_COLUMNS: tuple[str, ...] = (
    "파일명", "권리자", "제목", "저작자",
    *[f"C_{i}" for i in range(1, 11)], *[f"R_{i}" for i in range(1, 8)],
    *[f"이용허락기간_{p}{u}" for p in "SE" for u in "YMD"],
)
RIGHTS_LEDGER_COLUMNS: dict[str, str] = {
    "reproduction_right": "R_1", "performance_right": "R_2", "public_transmission_right": "R_3",
    "exhibition_right": "R_4", "distribution_right": "R_5", "rental_right": "R_6",
    "derivative_work_right": "R_7",
}
WORK_TYPE_LEDGER_COLUMNS: tuple[tuple[str, str], ...] = tuple(
    (f"C_{i}", WORK_TYPES[i - 1]) for i in range(1, 11))  # C_1 어문 … C_9 컴퓨터프로그램, C_10 기타
CHECKED_GLYPHS: frozenset[str] = frozenset({"v", "V", "■", "✓", "✔", "☑", "√"})
UNCHECKED_GLYPHS: frozenset[str] = frozenset({"□", "☐"})
LEDGER_MEDIA: dict[str, str] = {"어문": "text", "이미지": "image", "영상": "video"}
LEDGER_BUCKET: dict[str, str] = {"CCL": "ccl", "만료": "expired", "KOGL": "kogl", "기증": "donated",
                                 "만료-KOGL": "expired_kogl"}
# 원장 파일명 '<id> <제목>' — 앞의 생성 순번('00002_')과 뒤의 다운로드 중복 표시(' (1)')는 허용한다
# (장세영 2026-09-22 전달 5건: '00002_13387752 운수 자판기 (1).pdf' …). id 는 group(1).
LEDGER_ID_RE = re.compile(r"^(?:\d+_)?(\d{5,8})\s+(.*?)(?:\s*\(\d+\))?$", re.S)
LEDGER_MIN_ID_RATIO = 0.5  # id 를 읽은 행이 이보다 적으면 원장 형식이 바뀐 것 — 전부 na 인 gold 를 조용히 만들지 않고 실패시킨다

# ---------------------------------------------------------------- 사이트 설명 kind (generate_descriptions 와 동일 규칙)
LABEL_MIN_ROWS = 3  # 같은 설명이 같은 media×bucket×source 안에서 3행 이상이면 컬렉션 라벨
ATTRIB_ONLY_RE = re.compile(r"^.{1,40}의 (국외 )?(일러스트 작품|작품|사진 자료|자료)$")
URL_RE = re.compile(r"^https?://")
KOGL_SOURCE_PREFIX = "공공누리"
SITE_KINDS: tuple[str, ...] = ("none", "label", "plot", "background", "caption")

ANNOTATION_RE = re.compile(r"\(\s*(?:저작물\s*)?\d+\s*건\s*\)")  # (저작물 3 건) / (3건)
NAME_SPLIT_RE = re.compile(r"[\n,;·/]")
ZERO_WIDTH_RE = re.compile(r"[​-‏‪-‮⁠﻿]")
TAG_SEPARATOR = "|"
REJECTED_OUTCOMES: frozenset[str] = frozenset(
    {"rejected", "reject", "unusable", "fail", "failed", "excluded", "거절", "제외", "불가", "반려"})
DRAFT_STATUS_RANK: dict[str, int] = {"ok": 2, "needs_review": 1}
TIER_LEDGER = "A"
TIER_DRAFT = "A-검수(대기)"
TIER_VERIFIED = "A-검수"
TIER_FILE = "B"


# ================================================================ 공용 텍스트 도우미
def clean_text(value: Any) -> str:
    """NFC 정규화, zero-width 문자 제거, 공백 정리 (값 비교용 정규화는 평가기가 따로 한다)."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFC", str(value))
    text = ZERO_WIDTH_RE.sub("", text)
    return re.sub(r"\s+", " ", text).strip()


def tag_key(tag: str) -> str:
    """중복 제거용 태그 키 — 영숫자·한글만 남기고 casefold (평가기 norm_tag 와 같은 뜻)."""
    return "".join(ch for ch in tag if ch.isalnum()).casefold()


def dedupe(items: Iterable[str], key: Callable[[str], str] = tag_key) -> list[str]:
    """순서를 지키며 정규화 키 기준으로 중복을 없앤다."""
    seen: set[str] = set()
    out: list[str] = []
    for item in items:
        k = key(item)
        if k and k not in seen:
            seen.add(k)
            out.append(item)
    return out


def split_names(raw: Any) -> list[str]:
    """'양지웅 (저작물 3 건)\\n이영지' → ['양지웅', '이영지'] (주석 제거, 구분자 \\n , ; · / 로 분리)."""
    text = unicodedata.normalize("NFC", str(raw or ""))
    text = ANNOTATION_RE.sub(" ", text)
    parts = [clean_text(p) for p in NAME_SPLIT_RE.split(text)]
    return dedupe([p for p in parts if p], key=lambda s: re.sub(r"\s+", "", s).casefold())


def split_tags(raw: Any, sep: str = TAG_SEPARATOR) -> list[str]:
    """인덱스의 'a | b | c' 또는 목록 → ['a', 'b', 'c'] (앞의 '#' 제거, 중복 제거)."""
    items = raw if isinstance(raw, (list, tuple)) else str(raw or "").split(sep)
    cleaned = [clean_text(t).lstrip("#").strip() for t in items]
    return dedupe([t for t in cleaned if t])


def cell_str(value: Any) -> str:
    """openpyxl 셀 값 → 문자열 (2009 → '2009', 9.0 → '9', None → '')."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


# ================================================================ Gold 셀
def make_cell(value: Any, state: str, source: str, tier: str = "", kind: Optional[str] = None) -> dict[str, Any]:
    """§2 키 셀 — {"value", "state", "source", "tier", "kind"}."""
    return {"value": value, "state": state, "source": source, "tier": tier, "kind": kind}


def na_cell(reason: str, kind: Optional[str] = None) -> dict[str, Any]:
    return make_cell(None, STATE_NA, reason, "", kind)


def value_or_na(value: Any, source: str, reason: str, tier: str = "", kind: Optional[str] = None) -> dict[str, Any]:
    """빈 값(''/None/[])이면 na, 아니면 value."""
    if value is None or value == "" or value == []:
        return na_cell(reason, kind)
    return make_cell(value, STATE_VALUE, source, tier, kind)


# ================================================================ 입력 적재
@dataclass(frozen=True)
class SetRow:
    """id 목록의 한 행 — frozen csv 의 열이 없으면 빈 문자열 (인덱스/원장에서 보완한다)."""
    work_id: str
    media: str = ""
    bucket: str = ""
    source: str = ""
    work_file: str = ""
    title: str = ""
    contract_pdf: str = ""  # T-001: 시스템에 주는 계약서 PDF 경로 (file_name gold 의 기본 출처)


def _normalize_header(fields: Sequence[str]) -> list[str]:
    return [f.lstrip("﻿").strip() for f in fields]


def load_ids(path: Path) -> list[SetRow]:
    """frozen 목록(work_id 열 포함 csv) 또는 1열 id 파일(헤더 선택)을 읽는다. 중복 id 는 첫 행만 남긴다."""
    with open(path, encoding="utf-8-sig", newline="") as fh:
        rows = [r for r in csv.reader(fh) if r and any(c.strip() for c in r)]
    if not rows:
        return []
    header = _normalize_header(rows[0])
    if "work_id" in header:
        records = [dict(zip(header, r)) for r in rows[1:]]
        parsed = [SetRow(work_id=clean_text(d.get("work_id")), media=clean_text(d.get("media")),
                         bucket=clean_text(d.get("license_bucket")), source=clean_text(d.get("source")),
                         work_file=clean_text(d.get("work_file")), title=clean_text(d.get("title")),
                         contract_pdf=clean_text(d.get("contract_pdf")))
                  for d in records]
    else:
        first = [clean_text(r[0]).split(" ")[0] for r in rows]  # 1열 id (공백 뒤 제목이 붙어 있어도 됨)
        first = [w for w in first if w and not w.startswith("#")]
        if first and not first[0].isdigit():
            first = first[1:]  # 헤더 행(id 등)
        parsed = [SetRow(work_id=w) for w in first]
    out: list[SetRow] = []
    seen: set[str] = set()
    for row in parsed:
        if not row.work_id:
            continue
        if row.work_id in seen:
            log.warning("id 목록에 중복 id — 첫 행만 사용: %s", row.work_id)
            continue
        seen.add(row.work_id)
        out.append(row)
    log.info("id 목록 적재: %s → %d건", path, len(out))
    return out


def load_index(path: Path) -> dict[str, dict[str, str]]:
    """gt_merged_index.csv → {work_id: 행}. 전부 문자열, 빈칸은 '' (NaN 없음)."""
    df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    if "work_id" not in df.columns:
        raise ValueError(f"인덱스에 work_id 열이 없다: {path}")
    dup = df["work_id"].duplicated().sum()
    if dup:
        log.warning("인덱스에 중복 work_id %d건 — 첫 행만 사용", int(dup))
        df = df.drop_duplicates("work_id", keep="first")
    rows: dict[str, dict[str, str]] = {}
    for rec in df.to_dict("records"):
        rows[clean_text(rec["work_id"])] = {k: ("" if v is None else str(v)) for k, v in rec.items()}
    log.info("인덱스 적재: %s → %d행 × %d열", path, len(rows), len(df.columns))
    return rows


def load_ledger(path: Path, sheet: str = DEFAULT_LEDGER_SHEET) -> dict[str, dict[str, str]]:
    """원장 xlsx 시트 → {work_id: 행}. work_id 는 파일명 '<id> <제목>' 의 앞자리. 모르는 열은 로그로 알린다."""
    import openpyxl  # 무거운 의존성은 함수 안에서

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"원장에 시트 '{sheet}' 가 없다 (있는 시트: {wb.sheetnames})")
    rows_iter = wb[sheet].iter_rows(values_only=True)
    header = [cell_str(h) for h in next(rows_iter, ())]
    _report_ledger_columns(header)
    ledger: dict[str, dict[str, str]] = {}
    n_rows = unparsed = 0
    for raw in rows_iter:
        if not any(v is not None and cell_str(v) != "" for v in raw):
            continue
        n_rows += 1
        row = {h: cell_str(v) for h, v in zip(header, raw) if h}
        work_id = ledger_work_id(row.get("파일명", ""))
        if not work_id:
            unparsed += 1
            continue
        if work_id in ledger:
            log.warning("원장에 중복 id — 첫 행만 사용: %s", work_id)
            continue
        ledger[work_id] = row
    _check_ledger_parse_ratio(path, n_rows, unparsed)
    log.info("원장 적재: %s[%s] → %d행", path, sheet, len(ledger))
    return ledger


def ledger_work_id(file_stem: str) -> str:
    """원장 파일명 → work_id ('00002_13387752 운수 자판기 (1)' → '13387752'); 형식이 아니면 ''."""
    match = LEDGER_ID_RE.match(clean_text(file_stem))
    return match.group(1) if match else ""


def _check_ledger_parse_ratio(path: Path, n_rows: int, unparsed: int) -> None:
    """id 를 읽은 행이 0건이거나 절반 미만이면 원장 형식이 달라진 것 — 경고만 남기고 T-001 전부 na 가 되는 대신 실패시킨다."""
    matched = n_rows - unparsed
    if n_rows == 0 or matched == 0 or matched < n_rows * LEDGER_MIN_ID_RATIO:
        raise ValueError(f"원장 파일명에서 id 를 읽은 행이 {matched}/{n_rows}건 — 파일명 형식 '<id> <제목>' "
                         f"(앞 '00002_' 순번·뒤 ' (1)' 허용)인지 확인: {path}")
    if unparsed:
        log.warning("원장 파일명에서 id 를 읽지 못한 행 %d/%d건 (형식 '<id> <제목>' 아님)", unparsed, n_rows)


def _report_ledger_columns(header: Sequence[str]) -> None:
    present = set(h for h in header if h)
    unknown = sorted(present - LEDGER_KNOWN_COLUMNS)
    missing = [c for c in LEDGER_REQUIRED_COLUMNS if c not in present]
    if unknown:
        log.warning("원장에 모르는 열 %d개 (무시): %s", len(unknown), unknown)
    if missing:
        log.warning("원장에 필요한 열이 없다 (해당 키는 na): %s", missing)


def _iter_jsonl(path: Path) -> Iterable[dict[str, Any]]:
    with open(path, encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if line:
                yield json.loads(line)


def load_drafts(path: Path) -> dict[str, dict[str, Any]]:
    """drafts.jsonl → {work_id: 최선의 초안} (status ok > needs_review > 기타; 같은 등급이면 뒤의 행)."""
    best: dict[str, dict[str, Any]] = {}
    for rec in _iter_jsonl(path):
        work_id = clean_text(rec.get("work_id"))
        rank = DRAFT_STATUS_RANK.get(str(rec.get("status")), 0)
        if work_id and rank >= DRAFT_STATUS_RANK.get(str(best.get(work_id, {}).get("status")), -1):
            best[work_id] = rec
    log.info("생성 초안 적재: %s → %d세트 (ok %d)", path, len(best),
             sum(1 for r in best.values() if r.get("status") == "ok"))
    return best


def load_verified(path: Path) -> dict[str, dict[str, Any]]:
    """gold_descriptions.jsonl(검수 완료) → {work_id: 행}; 필드 description, hashtags, verification_outcome."""
    rows = {clean_text(r.get("work_id")): r for r in _iter_jsonl(path) if clean_text(r.get("work_id"))}
    log.info("검수 Gold 적재: %s → %d세트", path, len(rows))
    return rows


# ================================================================ 생성 Gold 선택
@dataclass(frozen=True)
class Generated:
    text: str
    hashtags: tuple[str, ...]
    tier: str
    source: str


@dataclass(frozen=True)
class GeneratedPick:
    gen: Optional[Generated]
    note: str  # gen 이 None 인 이유


def _pick_draft(work_id: str, drafts: Mapping[str, dict[str, Any]]) -> GeneratedPick:
    rec = drafts.get(work_id)
    if rec is None:
        return GeneratedPick(None, "생성 초안 없음")
    if rec.get("status") != "ok":
        return GeneratedPick(None, f"생성 초안 status={rec.get('status')}")
    text = clean_text(rec.get("draft_text"))
    if not text:
        return GeneratedPick(None, "생성 초안 본문 비어 있음")
    return GeneratedPick(Generated(text, tuple(split_tags(rec.get("draft_hashtags") or [])),
                                   TIER_DRAFT, "drafts.jsonl(ai_draft)"), "")


def _pick_verified(work_id: str, verified: Mapping[str, dict[str, Any]],
                   drafts: Mapping[str, dict[str, Any]]) -> GeneratedPick:
    rec = verified.get(work_id)
    if rec is None:  # 아직 검수되지 않은 세트 → 초안으로 보완 (tier 로 구분된다)
        pick = _pick_draft(work_id, drafts)
        if pick.gen is None:
            return pick
        return GeneratedPick(Generated(pick.gen.text, pick.gen.hashtags, TIER_DRAFT, "drafts.jsonl(미검수 보완)"), "")
    outcome = clean_text(rec.get("verification_outcome")).casefold()
    if outcome in REJECTED_OUTCOMES:
        return GeneratedPick(None, f"검수 거절(verification_outcome={outcome})")
    text = clean_text(rec.get("description"))
    if not text:
        return GeneratedPick(None, "검수 Gold 본문 비어 있음")
    return GeneratedPick(Generated(text, tuple(split_tags(rec.get("hashtags") or [])),
                                   TIER_VERIFIED, f"gold_descriptions.jsonl({outcome or 'verified'})"), "")


def select_generated(work_id: str, mode: str, drafts: Mapping[str, dict[str, Any]],
                     verified: Mapping[str, dict[str, Any]]) -> GeneratedPick:
    """--generated 모드에 따라 세트의 생성 Gold(설명+해시태그)를 고른다."""
    if mode == "none":
        return GeneratedPick(None, "--generated none")
    if mode == "verified":
        return _pick_verified(work_id, verified, drafts)
    return _pick_draft(work_id, drafts)


# ================================================================ 사이트 설명 kind
def classify_site_kinds(index: Mapping[str, Mapping[str, str]]) -> dict[str, str]:
    """전체 인덱스를 보고 세트별 사이트 설명 kind 를 정한다 (label 판정은 전체 컬렉션 기준이라 전체가 필요)."""
    tier_a = [(w, r) for w, r in index.items() if r.get("work_설명_tier") == "A" and r.get("work_설명", "").strip()]
    counts: Counter = Counter((r.get("media"), r.get("license_bucket"), r.get("source"), r.get("work_설명"))
                              for _, r in tier_a)
    kinds: dict[str, str] = {}
    for work_id, row in index.items():
        kinds[work_id] = _site_kind(row, counts)
    log.info("사이트 설명 kind 분포(인덱스 전체): %s", dict(sorted(Counter(kinds.values()).items())))
    return kinds


def _site_kind(row: Mapping[str, str], counts: Mapping[tuple, int]) -> str:
    text = row.get("work_설명", "").strip()
    if row.get("work_설명_tier") != "A" or not text or URL_RE.match(text) or ATTRIB_ONLY_RE.match(text):
        return "none"
    if counts.get((row.get("media"), row.get("license_bucket"), row.get("source"), row.get("work_설명")), 0) >= LABEL_MIN_ROWS:
        return "label"
    if row.get("media") == "video" and row.get("license_bucket") == "expired":
        return "plot"
    if row.get("media") == "video" and str(row.get("source", "")).startswith(KOGL_SOURCE_PREFIX):
        return "background"
    return "caption"


# ================================================================ SHA-256 캐시
class HashCache:
    """work_file 의 SHA-256 — dataset/_probe/file_sha256.jsonl 에 path+size+mtime 키로 누적한다."""

    def __init__(self, path: Path) -> None:
        self.path = path
        self._entries: dict[tuple[str, int, int], str] = {}
        if path.exists():
            for rec in _iter_jsonl(path):
                try:
                    self._entries[(str(rec["path"]), int(rec["size"]), int(rec["mtime"]))] = str(rec["sha256"])
                except (KeyError, TypeError, ValueError):
                    continue
        log.info("SHA-256 캐시: %s (%d건)", path, len(self._entries))

    def sha256(self, file_path: str) -> tuple[str, str]:
        """(hex, 'cache'|'computed'); 파일이 없으면 FileNotFoundError."""
        st = os.stat(file_path)
        key = (file_path, int(st.st_size), int(st.st_mtime))
        if key in self._entries:
            return self._entries[key], "cache"
        digest = sha256_file(file_path)
        self._entries[key] = digest
        self._append(key, digest)
        return digest, "computed"

    def _append(self, key: tuple[str, int, int], digest: str) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with open(self.path, "a", encoding="utf-8") as fh:
            fh.write(json.dumps({"path": key[0], "size": key[1], "mtime": key[2], "sha256": digest},
                                ensure_ascii=False) + "\n")


def sha256_file(file_path: str, chunk: int = 1 << 20) -> str:
    h = hashlib.sha256()
    with open(file_path, "rb") as fh:
        for block in iter(lambda: fh.read(chunk), b""):
            h.update(block)
    return h.hexdigest()


# ================================================================ T-001 키 (원장)
def ledger_keys(lrow: Optional[Mapping[str, str]]) -> dict[str, dict[str, Any]]:
    """키 1–16 을 원장 한 행에서 만든다. 행이 없으면 전부 na."""
    if lrow is None:
        return {k: na_cell("원장(ledger) 행 없음 — 계약서 미생성") for k in T001_KEYS}
    start, end = ledger_date(lrow, "S"), ledger_date(lrow, "E")
    out: dict[str, dict[str, Any]] = {
        "work_title": value_or_na(clean_text(lrow.get("제목")), "ledger:제목", "원장 제목 비어 있음", TIER_LEDGER),
        "file_name": ledger_file_name_cell(lrow),
        "work_type": work_type_cell(lrow),
        "creator": value_or_na(split_names(lrow.get("저작자")), "ledger:저작자", "원장 저작자 비어 있음", TIER_LEDGER),
        "copyright_holder": holder_cell(lrow.get("권리자")),
        "other_rights": na_cell("계약서에 기타 권리 항목 없음 — 전 계약서 상태 ③ (NULL 출력, 채점 제외)"),
        "license_start_date": value_or_na(start, "ledger:이용허락기간_SY/SM/SD", "원장 시작일 없음/형식 오류", TIER_LEDGER),
        "license_end_date": value_or_na(end, "ledger:이용허락기간_EY/EM/ED", "원장 종료일 없음/형식 오류", TIER_LEDGER),
        "contract_validity_period": value_or_na(f"{start}~{end}" if start and end else "",
                                                "ledger:이용허락기간 시작~종료", "시작일/종료일 중 하나 없음", TIER_LEDGER),
    }
    for key, col in RIGHTS_LEDGER_COLUMNS.items():
        out[key] = rights_cell(lrow.get(col), col)
    return out


def ledger_file_name_cell(lrow: Optional[Mapping[str, str]]) -> dict[str, Any]:
    """'<원장 파일명>.pdf' (§3). 장세영 쪽 명명 요청('<원문인덱스> <제목>.pdf')이 지켜지면 전달 파일명과 같다."""
    if lrow is None:
        return na_cell("원장(ledger) 행 없음 — 계약서 미생성")
    stem = clean_text(lrow.get("파일명"))
    return value_or_na(f"{stem}.pdf" if stem else "", "ledger:파일명+'.pdf'", "원장 파일명 비어 있음", TIER_LEDGER)


def contract_file_name_cell(row: SetRow, lrow: Optional[Mapping[str, str]], mode: str) -> dict[str, Any]:
    """T-001 file_name = 시스템에 실제로 주는 계약서 PDF 의 basename (§1 #2).

    auto  : id 목록의 contract_pdf 가 있으면 그 basename, 없으면 '<원장 파일명>.pdf'
    ids   : contract_pdf 만 (없으면 na)        ledger: '<원장 파일명>.pdf' 만
    전달 파일명이 원장과 다를 때('00002_13387752 운수 자판기 (1).pdf')도 gold 가 실제 파일명을 따르게 한다."""
    delivered = Path(row.contract_pdf).name if row.contract_pdf else ""
    if mode in ("auto", "ids") and delivered:
        return make_cell(delivered, STATE_VALUE, "ids:contract_pdf basename", TIER_FILE)
    if mode == "ids":
        return na_cell("id 목록에 contract_pdf 없음 (--file-name-from ids)")
    return ledger_file_name_cell(lrow)


def glyph_state(raw: Any) -> Optional[bool]:
    """체크 글리프 → True(허용), 빈 네모 → False(불허), 그 외 → None."""
    g = clean_text(raw)
    if g in CHECKED_GLYPHS:
        return True
    if g in UNCHECKED_GLYPHS:
        return False
    return None


def rights_cell(raw: Any, col: str) -> dict[str, Any]:
    checked = glyph_state(raw)
    if checked is None:
        return na_cell(f"원장 {col} 글리프 해석 불가: {clean_text(raw)!r}")
    return make_cell("허용" if checked else "불허", STATE_VALUE, f"ledger:{col}", TIER_LEDGER)


def work_type_cell(lrow: Mapping[str, str]) -> dict[str, Any]:
    """C_1..C_10 중 체크된 칸 → 폐쇄 집합 값. 둘 이상이면 첫 칸(경고), 없으면 na."""
    checked = [(col, name) for col, name in WORK_TYPE_LEDGER_COLUMNS if glyph_state(lrow.get(col)) is True]
    if not checked:
        return na_cell("원장 C_1..C_10 체크 없음")
    if len(checked) > 1:
        log.warning("원장 %s: 저작물 유형이 %d개 체크됨 → 첫 칸 사용 %s", lrow.get("파일명"), len(checked), checked)
    col, name = checked[0]
    source = f"ledger:{col}" + (f" (C_기타={clean_text(lrow.get('C_기타'))})" if col == "C_10" else "")
    return make_cell(name, STATE_VALUE, source, TIER_LEDGER)


def holder_cell(raw: Any) -> dict[str, Any]:
    """copyright_holder ← 권리자(새 템플릿의 이용자 당사자). '-'/빈값 → 상태 ③."""
    text = clean_text(raw)
    if text in ("", "-"):
        return na_cell("원장 권리자 '-'/빈값 — 계약서에 미인쇄 (상태 ③)")
    return make_cell(split_names(text), STATE_VALUE, "ledger:권리자 (이용자 당사자)", TIER_LEDGER)


def ledger_date(lrow: Mapping[str, str], which: str) -> str:
    """이용허락기간_{S|E}Y/M/D → 'YYYY-MM-DD'; 결손·범위 오류면 ''."""
    try:
        y, m, d = (int(float(lrow.get(f"이용허락기간_{which}{u}", ""))) for u in "YMD")
        return date(y, m, d).isoformat()
    except (TypeError, ValueError):
        return ""


# ================================================================ T-002 키 (인덱스 + 생성 Gold)
def index_field(irow: Mapping[str, str], col: str) -> tuple[str, Optional[str]]:
    """(값, 열이 없을 때의 사유). 인덱스 열이 없으면 na 로 처리하도록 사유를 돌려준다."""
    if col not in irow:
        return "", f"인덱스에 열 없음: {col}"
    return clean_text(irow[col]), None


def keywords_cell(irow: Mapping[str, str]) -> dict[str, Any]:
    content, _ = index_field(irow, "work_키워드_내용태그")  # 열이 없으면 work_키워드 로 내려간다
    if content:
        return make_cell(split_tags(content), STATE_VALUE, "index:work_키워드_내용태그",
                         irow.get("work_키워드_tier", ""))
    raw, missing_kw = index_field(irow, "work_키워드")
    if missing_kw:
        return na_cell(missing_kw)
    if irow.get("work_키워드_batch_level", "") == "True":
        return na_cell("배치 수준 컬렉션 태그 — 키워드 Gold 로 쓰지 않음 (review §8.3)")
    tags = split_tags(raw)
    if tags:
        return make_cell(tags, STATE_VALUE, "index:work_키워드", irow.get("work_키워드_tier", ""))
    if irow.get("work_키워드_tier", "") == "A":
        return make_cell([], STATE_EMPTY, "index:work_키워드 (수집됨, 비어 있음)", "A")
    return na_cell("사이트 키워드 미수집")


def hashtags_cell(row: SetRow, irow: Mapping[str, str], pick: GeneratedPick) -> dict[str, Any]:
    site_source = row.source or irow.get("source", "")
    raw, missing = index_field(irow, "work_해시태그")
    tier = irow.get("work_해시태그_tier", "")
    if not missing and site_source.startswith(KOGL_SOURCE_PREFIX):
        tags = split_tags(raw)
        if tags:
            return make_cell(tags, STATE_VALUE, "index:work_해시태그 (KOGL 원본)", tier, "site")
        if tier == "A":
            return make_cell([], STATE_EMPTY, "index:work_해시태그 (수집됨, 비어 있음)", tier, "site")
    if pick.gen is not None and pick.gen.hashtags:
        return make_cell(list(pick.gen.hashtags), STATE_VALUE, pick.gen.source, pick.gen.tier, "generated")
    why = "KOGL 원본 해시태그 없음" if site_source.startswith(KOGL_SOURCE_PREFIX) else "공유마당 — 사이트 해시태그 없음"
    return na_cell(f"{why}; 생성 Gold 없음 ({pick.note or '해시태그 비어 있음'})")


def file_format_cell(irow: Mapping[str, str]) -> dict[str, Any]:
    raw, missing = index_field(irow, "work_파일포맷")
    if missing:
        return na_cell(missing)
    return value_or_na(raw.upper().lstrip("."), "index:work_파일포맷", "파일포맷 미수집", irow.get("work_파일포맷_tier", ""))


def file_size_cell(irow: Mapping[str, str]) -> dict[str, Any]:
    raw, missing = index_field(irow, "work_파일크기")
    if missing:
        return na_cell(missing)
    try:
        size = int(float(raw.replace(",", ""))) if raw else None
    except ValueError:
        return na_cell(f"파일크기 해석 불가: {raw!r}")
    return value_or_na(size, "index:work_파일크기", "파일크기 미수집", irow.get("work_파일크기_tier", ""))


def resolution_cell(row: SetRow, irow: Mapping[str, str]) -> dict[str, Any]:
    media = row.media or irow.get("media", "")
    if media == "text":
        return na_cell("어문(text) 저작물 — 해상도 비적용")
    raw, missing = index_field(irow, "work_해상도")
    if missing:
        return na_cell(missing)
    return value_or_na(raw.replace(" ", ""), "index:work_해상도", "해상도 미수집", irow.get("work_해상도_tier", ""))


def description_cell(irow: Mapping[str, str], pick: GeneratedPick, site_kind: str) -> dict[str, Any]:
    """caption → 사이트 Gold; 그 외 kind → 생성 Gold 또는 na (kind 는 na 에도 남겨 사유를 알 수 있게 한다)."""
    if site_kind == "caption":
        return make_cell(clean_text(irow.get("work_설명")), STATE_VALUE, "index:work_설명 (사이트 캡션)",
                         irow.get("work_설명_tier", ""), "caption")
    if pick.gen is not None:
        return make_cell(pick.gen.text, STATE_VALUE, pick.gen.source, pick.gen.tier, "generated")
    why = "사이트 설명 없음" if site_kind == "none" else f"사이트 설명 kind={site_kind} — 코사인 채점 불가"
    return na_cell(f"{why}; 생성 Gold 없음 ({pick.note})", None if site_kind == "none" else site_kind)


def hash_cell(work_file: str, hasher: Optional[HashCache]) -> dict[str, Any]:
    if hasher is None:
        return na_cell("미계산 (--hash 필요)")
    if not work_file:
        return na_cell("work_file 경로 없음")
    try:
        digest, how = hasher.sha256(work_file)
    except FileNotFoundError:
        return na_cell(f"파일 없음: {work_file}")
    except OSError as exc:
        return na_cell(f"파일 접근 실패: {exc}")
    return make_cell(digest, STATE_VALUE, f"sha256(work_file) [{how}]", TIER_FILE)


def work_file_name_cell(work_file: str) -> dict[str, Any]:
    name = Path(work_file).name if work_file else ""
    return value_or_na(name, "work_file basename", "work_file 경로 없음", TIER_FILE)


def work_keys(row: SetRow, irow: Optional[Mapping[str, str]], pick: GeneratedPick, site_kind: str,
              hasher: Optional[HashCache]) -> dict[str, dict[str, Any]]:
    """키 17–23 (+ T-002 file_name) 을 인덱스 행·생성 Gold 에서 만든다."""
    if irow is None:
        out = {k: na_cell("인덱스에 work_id 없음") for k in T002_KEYS}
        out["file_name"] = work_file_name_cell(row.work_file)
        return out
    work_file = row.work_file or irow.get("work_file", "")
    return {
        "file_name": work_file_name_cell(work_file),
        "keywords": keywords_cell(irow),
        "hashtags": hashtags_cell(row, irow, pick),
        "file_format": file_format_cell(irow),
        "file_size": file_size_cell(irow),
        "resolution": resolution_cell(row, irow),
        "description": description_cell(irow, pick, site_kind),
        "file_hash": hash_cell(work_file, hasher),
    }


# ================================================================ 레코드 조립
def resolve_meta(row: SetRow, irow: Optional[Mapping[str, str]], lrow: Optional[Mapping[str, str]]) -> tuple[str, str, str]:
    """(media, bucket, source) — id 목록 → 인덱스 → 원장 순으로 보완."""
    irow = irow or {}
    lrow = lrow or {}
    media = row.media or irow.get("media", "") or LEDGER_MEDIA.get(lrow.get("저작물형태", ""), "")
    bucket = row.bucket or irow.get("license_bucket", "") or LEDGER_BUCKET.get(lrow.get("공유형태", ""), "")
    source = row.source or irow.get("source", "")
    return media, bucket, source


def build_record(test: str, row: SetRow, irow: Optional[Mapping[str, str]], lrow: Optional[Mapping[str, str]],
                 pick: GeneratedPick, site_kind: str, hasher: Optional[HashCache],
                 file_name_from: str = "auto") -> dict[str, Any]:
    """세트 한 건의 §2 레코드. 23키 모두 KEYS 순서로 들어간다."""
    media, bucket, source = resolve_meta(row, irow, lrow)
    if test == "T-001":
        keys = ledger_keys(lrow)
        keys["file_name"] = contract_file_name_cell(row, lrow, file_name_from)
        keys.update({k: na_cell("T-002 전용 키 — T-001 비채점") for k in T002_KEYS})
    else:
        keys = {k: na_cell("T-001 전용 키 — T-002 비채점") for k in T001_KEYS}
        keys.update(work_keys(SetRow(row.work_id, media, bucket, source, row.work_file, row.title), irow, pick, site_kind, hasher))
    return {"set_id": row.work_id, "test": test, "media": media, "bucket": bucket, "source": source,
            "keys": {k: keys[k] for k in KEY_NAMES}}


def build_gold(test: str, ids_path: Path, index_path: Path, ledger_path: Optional[Path],
               drafts_path: Optional[Path], verified_path: Optional[Path], generated_mode: str,
               hash_flag: bool, hash_cache_path: Path, ledger_sheet: str = DEFAULT_LEDGER_SHEET,
               file_name_from: str = "auto") -> list[dict[str, Any]]:
    """§3 — frozen 목록의 세트마다 Gold 레코드를 만든다 (id 목록 순서 유지, 결정적)."""
    if test not in TESTS:
        raise ValueError(f"test 는 {TESTS} 중 하나: {test!r}")
    if generated_mode not in GENERATED_MODES:
        raise ValueError(f"generated_mode 는 {GENERATED_MODES} 중 하나: {generated_mode!r}")
    if file_name_from not in FILE_NAME_MODES:
        raise ValueError(f"file_name_from 은 {FILE_NAME_MODES} 중 하나: {file_name_from!r}")
    rows = load_ids(Path(ids_path))
    index = load_index(Path(index_path))
    ledger = load_ledger(Path(ledger_path), ledger_sheet) if test == "T-001" and ledger_path else {}
    drafts, verified = _load_generated(test, drafts_path, verified_path, generated_mode)
    kinds = classify_site_kinds(index) if test == "T-002" else {}
    hasher = HashCache(Path(hash_cache_path)) if (hash_flag and test == "T-002") else None
    if test == "T-001" and not ledger:
        log.warning("T-001 인데 원장이 없다 — 키 1–16 전부 na")
    records: list[dict[str, Any]] = []
    for row in rows:
        irow = index.get(row.work_id)
        if irow is None:
            log.warning("인덱스에 없는 work_id: %s", row.work_id)
        pick = select_generated(row.work_id, generated_mode, drafts, verified) if test == "T-002" else GeneratedPick(None, "T-001")
        records.append(build_record(test, row, irow, ledger.get(row.work_id), pick, kinds.get(row.work_id, "none"), hasher,
                                    file_name_from))
    if test == "T-001":
        _log_file_name_disagreements(rows, ledger)
    _log_summary(test, records, ledger, rows)
    return records


def _log_file_name_disagreements(rows: Sequence[SetRow], ledger: Mapping[str, Mapping[str, str]]) -> None:
    """id 목록의 contract_pdf basename 과 '<원장 파일명>.pdf' 가 다른 세트를 알린다 (전달 파일명 ≠ 원장 명명)."""
    pairs = [(Path(r.contract_pdf).name, f"{clean_text(ledger[r.work_id].get('파일명'))}.pdf")
             for r in rows if r.contract_pdf and r.work_id in ledger]
    diff = [(a, b) for a, b in pairs if a != b]
    if diff:
        log.warning("id 목록(contract_pdf)과 원장 파일명이 다른 세트 %d/%d건 — 예: %r ≠ %r (--file-name-from 기준으로 결정)",
                    len(diff), len(pairs), diff[0][0], diff[0][1])


def _load_generated(test: str, drafts_path: Optional[Path], verified_path: Optional[Path],
                    mode: str) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    """T-002 에서만 초안/검수 Gold 를 읽는다. verified 모드인데 검수 파일이 없으면 오류."""
    if test != "T-002" or mode == "none":
        return {}, {}
    drafts = load_drafts(Path(drafts_path)) if drafts_path and Path(drafts_path).exists() else {}
    if drafts_path and not drafts:
        log.warning("생성 초안 파일이 없거나 비어 있다: %s", drafts_path)
    if mode == "draft":
        return drafts, {}
    vpath = Path(verified_path) if verified_path else (Path(drafts_path).parent / VERIFIED_BASENAME if drafts_path else None)
    if vpath is None or not vpath.exists():
        raise FileNotFoundError(f"--generated verified 인데 검수 Gold 파일이 없다: {vpath}")
    return drafts, load_verified(vpath)


def _log_summary(test: str, records: Sequence[dict[str, Any]], ledger: Mapping[str, Any], rows: Sequence[SetRow]) -> None:
    if test == "T-001":
        matched = sum(1 for r in rows if r.work_id in ledger)
        log.info("원장 매칭: %d/%d 세트 (나머지는 키 1–16 na)", matched, len(rows))
    else:
        kinds = Counter(r["keys"]["description"]["kind"] or "none" for r in records)
        log.info("설명 Gold kind 분포(상태 포함): %s", dict(sorted(kinds.items())))
    log.info("Gold 레코드 %d건 생성 (%s)", len(records), test)


# ================================================================ 커버리지 표
def coverage(records: Sequence[dict[str, Any]], keys: Sequence[str]) -> dict[str, dict[str, Counter]]:
    """{cell(media/bucket): {key: Counter(state)}}."""
    cov: dict[str, dict[str, Counter]] = {}
    for rec in records:
        cell = f"{rec['media'] or '?'}/{rec['bucket'] or '?'}"
        per_key = cov.setdefault(cell, {})
        for k in keys:
            per_key.setdefault(k, Counter())[rec["keys"][k]["state"]] += 1
    return cov


def _fmt_counts(cnt: Counter) -> str:
    return f"{cnt.get(STATE_VALUE, 0)}/{cnt.get(STATE_EMPTY, 0)}/{cnt.get(STATE_NA, 0)}"


def format_coverage(cov: Mapping[str, Mapping[str, Counter]], keys: Sequence[str]) -> str:
    """행=키, 열=셀(+합계), 칸='value/empty/na' 건수."""
    cells = sorted(cov)
    header = ["key"] + cells + ["합계"]
    table: list[list[str]] = []
    for k in keys:
        total: Counter = Counter()
        line = [k]
        for c in cells:
            cnt = cov[c].get(k, Counter())
            total.update(cnt)
            line.append(_fmt_counts(cnt))
        line.append(_fmt_counts(total))
        table.append(line)
    widths = [max(len(r[i]) for r in [header] + table) for i in range(len(header))]

    def fmt_row(r: Sequence[str]) -> str:
        return "  ".join(v.ljust(widths[i]) for i, v in enumerate(r))

    lines = ["셀 × 키 커버리지 (value/empty/na 건수)", fmt_row(header), fmt_row(["-" * w for w in widths])]
    lines += [fmt_row(r) for r in table]
    return "\n".join(lines)


def write_jsonl(records: Sequence[dict[str, Any]], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as fh:
        for rec in records:
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
    log.info("Gold 저장: %s (%d행)", out_path, len(records))


# ================================================================ CLI
def build_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(prog="python -m api.module.evaluation.v4.gold23",
                                 description="v4 23키 Gold 빌더 (M1 spec §2–§3)")
    ap.add_argument("--test", required=True, choices=TESTS)
    ap.add_argument("--ids", required=True, type=Path, help="frozen id 목록 csv 또는 1열 id 파일")
    ap.add_argument("--out", required=True, type=Path, help="출력 gold23 jsonl")
    ap.add_argument("--index", type=Path, default=DEFAULT_INDEX)
    ap.add_argument("--ledger", type=Path, default=DEFAULT_LEDGER, help="계약서 원장 xlsx (T-001)")
    ap.add_argument("--ledger-sheet", default=DEFAULT_LEDGER_SHEET)
    ap.add_argument("--drafts", type=Path, default=DEFAULT_DRAFTS, help="생성 설명 초안 jsonl (T-002)")
    ap.add_argument("--verified", type=Path, default=None,
                    help=f"검수 Gold jsonl (기본: drafts 와 같은 폴더의 {VERIFIED_BASENAME})")
    ap.add_argument("--generated", choices=GENERATED_MODES, default="draft")
    ap.add_argument("--file-name-from", choices=FILE_NAME_MODES, default="auto",
                    help="T-001 file_name gold 출처: auto(id 목록 contract_pdf 의 basename, 없으면 원장) | ids | ledger")
    ap.add_argument("--hash", action="store_true", help="work_file SHA-256 계산 (느림 — 캐시 사용)")
    ap.add_argument("--hash-cache", type=Path, default=DEFAULT_HASH_CACHE)
    ap.add_argument("--log-level", default="INFO")
    return ap


def configure_logging(level_name: str) -> None:
    """api 패키지가 import 시 root 로거를 basicConfig(WARNING) 로 잡아 두므로 이 모듈 로거만 따로 설정한다."""
    log.setLevel(getattr(logging, level_name.upper(), logging.INFO))
    if not log.handlers:
        handler = logging.StreamHandler(sys.stderr)
        handler.setFormatter(logging.Formatter("%(levelname)s %(name)s: %(message)s"))
        log.addHandler(handler)
    log.propagate = False


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    configure_logging(str(args.log_level))
    records = build_gold(args.test, args.ids, args.index, args.ledger, args.drafts, args.verified,
                         args.generated, args.hash, args.hash_cache, args.ledger_sheet, file_name_from=args.file_name_from)
    write_jsonl(records, args.out)
    keys = T001_KEYS if args.test == "T-001" else ("file_name",) + T002_KEYS
    print(format_coverage(coverage(records, keys), keys))
    return 0


if __name__ == "__main__":
    sys.exit(main())
