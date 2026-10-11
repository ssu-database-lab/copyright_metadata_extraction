"""영상 원본의 기술 메타데이터(가로·세로·fps·재생시간)를 ffprobe 로 실측해 정답 색인에 기록한다.

왜 실측인가 (2026-10-02 확인)
------------------------------
영상 해상도를 '외부에서' 받을 곳이 사실상 없다.

  - KOGL(공공누리): 붙임1 75열, dataset/manifest.xlsx 36열 어디에도 크기·해상도·재생시간
    열이 없다. 상세페이지는 "재생 영상은 해상도가 720p로 제공되며, 다운로드 영상은
    해상도에 다소 차이가 있을 수 있습니다" 라고만 적는다 — 실측해 보면 다운로드본
    500건 중 4K 456 · 1080p 38 · 720p 3건이다. 게시 문구는 플레이어 기준이다.
  - 공유마당: 상세페이지가 '영상저작물 화질'·'재생시간'을 게시하지만 만료 499/499,
    기증 251/498, CCL 12/1003 에만 있다. (재생시간 라벨이 영상에도 '음원저작물
    재생시간' 으로 붙어 있다 — 사이트 쪽 오기.)
  - 무하유 공유마당 엑셀(354k)의 영상 시트에는 재생시간·파일크기만 있고 해상도는 없다.

티어
----
실측값은 **B(파일 산출)** 다. 추출기(clip_extraction/technical_metadata.py)도 같은 파일을
ffprobe 로 읽으므로, 이 값과의 비교는 정확도가 아니라 "값이 업로드 → 추출 → 통합 →
직렬화를 거쳐 살아남는가" 의 검증이다. 사이트 게시값(A)이 있으면 그쪽이 외부 정답이고,
이 스크립트는 두 값이 실제로 맞는지도 대조한다.

산출
----
  1) 캐시   dataset/_probe/video_ffprobe.jsonl          파일 1개 = 1줄. (경로·크기·mtime) 가
                                                        같으면 재실행 시 다시 재지 않는다.
  2) KOGL   E:/kogl_originals/originals_index.xlsx     영상 행의 width·height·megapixels 를
                                                        채우고 fps·duration_sec·video_codec·
                                                        dims_source 열을 더한다. 이미지 행은
                                                        PIL 로 재측정해 기존값과 대조만 한다.
  3) 공유마당 {E:/gongu_dataset, D:/copyright_dataset_metadata}/video/{expired,donated,ccl}/
            gongu_index.xlsx                           측정열 + 사이트 게시값 + 일치 여부 열
  4) 대조표 D:/copyright_dataset_metadata/video/site_vs_ffprobe.csv

엑셀을 고치기 전에 원본을 <root>/_backup_YYYYMMDD/ 에 복사해 둔다.

사용
----
  python -m api.module.dataset_builder.measure_video_technical            # 측정 + 색인 기록
  python -m api.module.dataset_builder.measure_video_technical --probe-only
"""
from __future__ import annotations

import argparse
import csv
import datetime as _dt
import json
import os
import re
import shutil
import subprocess
import sys
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

KOGL_ROOT = Path("/mnt/e/kogl_originals")
KOGL_INDEX = KOGL_ROOT / "originals_index.xlsx"
GONGU_ROOTS = (Path("/mnt/e/gongu_dataset"), Path("/mnt/d/copyright_dataset_metadata"))
META_ROOT = Path("/mnt/d/copyright_dataset_metadata")      # detail_meta.jsonl 이 있는 쪽
BUCKETS = ("expired", "donated", "ccl")
REPO = Path(__file__).resolve().parents[3]
CACHE = REPO / "dataset" / "_probe" / "video_ffprobe.jsonl"
REPORT = META_ROOT / "video" / "site_vs_ffprobe.csv"

TODAY = _dt.date.today().isoformat()
PROBE_SOURCE = f"ffprobe 실측 {TODAY}"

# 사이트 해상도 값은 양방향 제어문자로 감싸여 있다 ('‪723 x 1680‬'). build_ground_truth 와 같은 처리.
_BIDI = dict.fromkeys(range(0x202A, 0x202F), None)


# ---------------------------------------------------------------- probe
def _rate(s) -> float | None:
    try:
        num, den = str(s).split("/", 1)
        num, den = float(num), float(den)
    except (ValueError, TypeError):
        return None
    if num <= 0 or den <= 0:
        return None
    r = num / den
    return r if 0 < r <= 240 else None       # WMV 의 1000/1 같은 시간축 값은 버린다


def _run(cmd, timeout=180):
    return subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)


def probe(path: str) -> dict:
    """첫 번째 '진짜' 영상 스트림(앨범아트 제외)의 기술 메타데이터."""
    out = {"path": path, "status": "ok"}
    try:
        st = os.stat(path)
        out["size"], out["mtime"] = st.st_size, int(st.st_mtime)
    except OSError as e:
        return {**out, "status": "missing", "error": str(e)}

    cmd = ["ffprobe", "-v", "error", "-show_entries",
           "stream=index,codec_type,codec_name,width,height,r_frame_rate,avg_frame_rate,"
           "duration,sample_aspect_ratio:stream_tags=rotate:stream_side_data=rotation:"
           "stream_disposition=attached_pic:format=duration,format_name",
           "-of", "json", path]
    try:
        r = _run(cmd)
        data = json.loads(r.stdout or "{}")
    except (subprocess.TimeoutExpired, json.JSONDecodeError) as e:
        return {**out, "status": "probe_failed", "error": f"{type(e).__name__}"}

    if not data.get("format"):
        # 컨테이너 자체를 못 읽음 — 예: moov atom 이 없는 잘린 .mov (다운로드 중단 흔적)
        err = (r.stderr or "").strip().splitlines()
        return {**out, "status": "unreadable", "error": err[0][-160:] if err else "no format"}

    vids = [s for s in data.get("streams") or []
            if s.get("codec_type") == "video"
            and not (s.get("disposition") or {}).get("attached_pic")]
    fmt = data.get("format") or {}
    out["format_name"] = fmt.get("format_name")
    if not vids:
        return {**out, "status": "no_video_stream"}
    v = vids[0]

    out["width"], out["height"] = v.get("width"), v.get("height")
    out["video_codec"] = v.get("codec_name")
    out["sar"] = v.get("sample_aspect_ratio")
    rot = None
    for sd in v.get("side_data_list") or []:
        if "rotation" in sd:
            rot = sd["rotation"]
    rot = rot if rot is not None else (v.get("tags") or {}).get("rotate")
    out["rotation"] = int(float(rot)) if rot not in (None, "") else 0

    # avg_frame_rate 가 실제 평균. r_frame_rate 는 WMV 에서 시간축(1000/1)이 나오기도 한다.
    fps = _rate(v.get("avg_frame_rate")) or _rate(v.get("r_frame_rate"))
    out["fps"] = round(fps, 3) if fps else None

    # 재생시간: 컨테이너 → 스트림 → (SWF 처럼 둘 다 없으면) 패킷 수 ÷ fps
    dur, how = None, None
    for src, val in (("format", fmt.get("duration")), ("stream", v.get("duration"))):
        try:
            if val not in (None, "N/A") and float(val) > 0:
                dur, how = float(val), src
                break
        except ValueError:
            pass
    if dur is None and fps:
        try:
            rr = _run(["ffprobe", "-v", "error", "-select_streams", f"{v.get('index', 0)}",
                       "-count_packets", "-show_entries", "stream=nb_read_packets",
                       "-of", "csv=p=0", path], timeout=600)
            n = int((rr.stdout or "0").strip().split(",")[0] or 0)
            if n > 0:
                dur, how = n / fps, "packets/fps"
        except (subprocess.TimeoutExpired, ValueError):
            pass
    out["duration_sec"] = round(dur, 3) if dur else None
    out["duration_method"] = how
    if not (out["width"] and out["height"]):
        out["status"] = "no_dimensions"
    return out


def load_cache() -> dict:
    cache = {}
    if CACHE.exists():
        for line in CACHE.open(encoding="utf-8"):
            try:
                r = json.loads(line)
                cache[r["path"]] = r
            except (json.JSONDecodeError, KeyError):
                continue
    return cache


def measure(paths, workers=8) -> dict:
    """paths 를 모두 잰다. 캐시가 (크기·mtime) 까지 같으면 재사용."""
    cache = load_cache()
    todo = []
    for p in sorted(set(paths)):
        c = cache.get(p)
        try:
            st = os.stat(p)
            fresh = c and c.get("size") == st.st_size and c.get("mtime") == int(st.st_mtime)
        except OSError:
            fresh = False
        if not fresh:
            todo.append(p)
    print(f"  측정 대상 {len(set(paths)):,}개 — 캐시 재사용 {len(set(paths)) - len(todo):,}, 새로 측정 {len(todo):,}",
          flush=True)
    if todo:
        CACHE.parent.mkdir(parents=True, exist_ok=True)
        with ThreadPoolExecutor(workers) as ex, CACHE.open("a", encoding="utf-8") as f:
            for i, r in enumerate(ex.map(probe, todo), 1):
                cache[r["path"]] = r
                f.write(json.dumps(r, ensure_ascii=False) + "\n")
                if i % 200 == 0:
                    print(f"    {i:,}/{len(todo):,}", flush=True)
    return cache


def load_measurements() -> dict:
    """다른 빌더용: {파일경로: 측정 레코드}. 측정을 새로 하지 않는다."""
    return load_cache()


def resolution(m: dict | None) -> str | None:
    if m and m.get("width") and m.get("height"):
        return f"{m['width']}x{m['height']}"
    return None


# ---------------------------------------------------------------- 사이트 게시값
def norm_site_resolution(s) -> str | None:
    t = unicodedata.normalize("NFKC", str(s or "")).translate(_BIDI).strip()
    m = re.match(r"^(\d+)\s*[xX×*]\s*(\d+)$", t)
    return f"{m.group(1)}x{m.group(2)}" if m else None


def site_seconds(s) -> int | None:
    t = unicodedata.normalize("NFKC", str(s or "")).translate(_BIDI).strip()
    m = re.match(r"^(\d+):(\d{1,2}):(\d{1,2})$", t) or re.match(r"^()(\d+):(\d{1,2})$", t)
    if not m:
        return None
    h, mi, se = (int(x) if x else 0 for x in m.groups())
    return h * 3600 + mi * 60 + se


def load_site_values(bucket: str) -> dict:
    """detail_meta.jsonl → {wrtSn: (해상도, 재생시간 원문, 화면비율)}"""
    p = META_ROOT / "video" / bucket / "detail_meta.jsonl"
    out = {}
    if not p.exists():
        return out
    for line in p.open(encoding="utf-8"):
        try:
            r = json.loads(line)
        except json.JSONDecodeError:
            continue
        if not r.get("ok"):
            continue
        f = r.get("fields") or {}
        out[str(r.get("wrtSn"))] = (
            norm_site_resolution(f.get("영상저작물 화질")),
            f.get("영상저작물 재생시간") or f.get("음원저작물 재생시간"),
            f.get("영상저작물 화면비율"),
        )
    return out


def compare_resolution(site: str | None, measured: str | None) -> str:
    """exact / swapped / same_aspect(배율만 다름) / differs / no_site / no_measure"""
    if not site:
        return "no_site"
    if not measured:
        return "no_measure"
    if site == measured:
        return "exact"
    sw, sh = map(int, site.split("x"))
    mw, mh = map(int, measured.split("x"))
    if (sw, sh) == (mh, mw):
        return "swapped"
    if abs(sw / sh - mw / mh) < 0.01:
        return "same_aspect"
    return "differs"


# ---------------------------------------------------------------- 엑셀 기록
def _backup(path: Path, root: Path):
    bdir = root / f"_backup_{TODAY.replace('-', '')}"
    rel = path.relative_to(root)
    dst = bdir / rel
    if not dst.exists():                       # 같은 날 두 번 돌려도 '수정 전' 원본을 지킨다
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, dst)
    return dst


def _ensure_cols(ws, names):
    hdr = {c.value: c.column for c in ws[1] if c.value is not None}
    for n in names:
        if n not in hdr:
            col = ws.max_column + 1
            ws.cell(row=1, column=col, value=n)
            hdr[n] = col
    return hdr


def update_kogl_index(cache: dict) -> dict:
    import openpyxl
    from PIL import Image

    _backup(KOGL_INDEX, KOGL_ROOT)
    wb = openpyxl.load_workbook(KOGL_INDEX)
    ws = wb["originals_index"]
    hdr = _ensure_cols(ws, ["fps", "duration_sec", "video_codec", "dims_source"])
    stats = {"video_filled": 0, "video_failed": 0, "image_verified": 0, "image_mismatch": []}

    for row in range(2, ws.max_row + 1):
        cls = ws.cell(row, hdr["분류"]).value
        saved = ws.cell(row, hdr["saved_file"]).value
        if not saved:
            continue
        if cls == "영상":
            m = cache.get(str(KOGL_ROOT / "영상" / saved))
            if not (m and m.get("status") == "ok"):
                stats["video_failed"] += 1
                continue
            w, h = m["width"], m["height"]
            ws.cell(row, hdr["width"], w)
            ws.cell(row, hdr["height"], h)
            ws.cell(row, hdr["megapixels"], round(w * h / 1e6, 2))
            ws.cell(row, hdr["fps"], m.get("fps"))
            ws.cell(row, hdr["duration_sec"], m.get("duration_sec"))
            ws.cell(row, hdr["video_codec"], m.get("video_codec"))
            ws.cell(row, hdr["dims_source"], PROBE_SOURCE)
            stats["video_filled"] += 1
        elif cls == "이미지":
            # 기존 width/height 를 만든 스크립트가 저장소에 없다 — 다시 재서 출처를 복원한다.
            p = KOGL_ROOT / "이미지" / saved
            try:
                with Image.open(p) as im:
                    pw, ph = im.size
            except Exception as e:                       # noqa: BLE001
                stats["image_mismatch"].append((saved, f"열기 실패 {type(e).__name__}"))
                continue
            ow, oh = ws.cell(row, hdr["width"]).value, ws.cell(row, hdr["height"]).value
            if ow and oh and (int(ow), int(oh)) == (pw, ph):
                ws.cell(row, hdr["dims_source"], f"PIL 실측 재확인 {TODAY}")
                stats["image_verified"] += 1
            else:
                stats["image_mismatch"].append((saved, f"색인 {ow}x{oh} ≠ 파일 {pw}x{ph}"))
    wb.save(KOGL_INDEX)
    return stats


GONGU_COLS = ["probe_file", "width", "height", "fps", "duration_sec", "video_codec",
              "probe_status", "probe_all_files", "site_resolution", "site_duration",
              "resolution_match", "duration_diff_sec"]


def gongu_rows(bucket: str):
    """gongu_index.xlsx 의 수집 성공 행 → (wrtSn, 제목, [파일경로...])"""
    import openpyxl
    p = META_ROOT / "video" / bucket / "gongu_index.xlsx"
    wb = openpyxl.load_workbook(p, read_only=True)
    ws = wb["gongu_index"]
    it = ws.iter_rows(values_only=True)
    hdr = {n: i for i, n in enumerate(next(it))}
    out = []
    for r in it:
        paths = [s.strip() for s in str(r[hdr["saved_paths"]] or "").split(";") if s.strip()]
        if r[hdr["status"]] == "ok" and paths:
            out.append((str(r[hdr["wrtSn"]]), r[hdr["제목"]], paths))
    wb.close()
    return out


def _fmt_file(m):
    if not m or m.get("status") != "ok":
        return f"{os.path.basename(m['path']) if m else '?'}: {m.get('status') if m else 'missing'}"
    return (f"{os.path.basename(m['path'])}: {resolution(m)} @{m.get('fps')}fps, "
            f"{m.get('duration_sec')}s ({m.get('video_codec')})")


def update_gongu_indexes(cache: dict) -> list:
    import openpyxl
    report = []
    for bucket in BUCKETS:
        site = load_site_values(bucket)
        rows = {wrt: (title, paths) for wrt, title, paths in gongu_rows(bucket)}
        computed = {}
        for wrt, (title, paths) in rows.items():
            ms = [cache.get(p) for p in paths]
            prim = ms[0]
            s_res, s_dur_raw, s_ar = site.get(wrt, (None, None, None))
            p_res = resolution(prim)
            verdict = compare_resolution(s_res, p_res)
            any_exact = next((os.path.basename(m["path"]) for m in ms if s_res and resolution(m) == s_res), None)
            s_dur = site_seconds(s_dur_raw)
            d_diff = (round(prim["duration_sec"] - s_dur, 2)
                      if (s_dur is not None and prim and prim.get("duration_sec")) else None)
            computed[wrt] = {
                "probe_file": os.path.basename(paths[0]),
                "width": prim.get("width") if prim else None,
                "height": prim.get("height") if prim else None,
                "fps": prim.get("fps") if prim else None,
                "duration_sec": prim.get("duration_sec") if prim else None,
                "video_codec": prim.get("video_codec") if prim else None,
                "probe_status": prim.get("status") if prim else "missing",
                "probe_all_files": " ; ".join(_fmt_file(m) for m in ms),
                "site_resolution": s_res,
                "site_duration": s_dur_raw,
                "resolution_match": verdict,
                "duration_diff_sec": d_diff,
            }
            report.append({
                "bucket": bucket, "wrtSn": wrt, "title": title, "n_files": len(paths),
                "primary_file": os.path.basename(paths[0]),
                "site_resolution": s_res, "measured_resolution": p_res, "resolution_match": verdict,
                "site_matches_other_file": any_exact if (verdict != "exact" and any_exact) else "",
                "site_aspect": s_ar,
                "site_duration": s_dur_raw, "site_duration_sec": s_dur,
                "measured_duration_sec": prim.get("duration_sec") if prim else None,
                "duration_method": prim.get("duration_method") if prim else None,
                "duration_diff_sec": d_diff,
                "measured_fps": prim.get("fps") if prim else None,
                "probe_status": prim.get("status") if prim else "missing",
            })

        for root in GONGU_ROOTS:
            p = root / "video" / bucket / "gongu_index.xlsx"
            if not p.exists():
                continue
            _backup(p, root)
            wb = openpyxl.load_workbook(p)
            ws = wb["gongu_index"]
            hdr = _ensure_cols(ws, GONGU_COLS)
            for row in range(2, ws.max_row + 1):
                vals = computed.get(str(ws.cell(row, hdr["wrtSn"]).value))
                if not vals:
                    continue
                for k, v in vals.items():
                    ws.cell(row, hdr[k], v)
            wb.save(p)
            print(f"  기록: {p}  ({len(computed):,}행)", flush=True)

    REPORT.parent.mkdir(parents=True, exist_ok=True)
    with REPORT.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(report[0]))
        w.writeheader()
        w.writerows(report)
    return report


# ---------------------------------------------------------------- 매니페스트 보강
MANIFEST_COLS = ["gt_width", "gt_height", "gt_resolution", "gt_fps", "gt_duration_sec",
                 "gt_resolution_source", "gt_resolution_site", "gt_duration_site_sec",
                 "gt_resolution_site_matches_file"]


def video_technical_columns(work_path: str, media: str, license_bucket: str, set_id: str,
                            cache: dict, site_by_bucket: dict) -> dict:
    """contract_work_manifest 한 행에 붙일 영상 기술 정답.

    측정 대상은 ZIP 에 실제로 실리는 work_path 다. 공유마당 정답(ground_truth.jsonl)의
    file 과 다른 변형(swf↔mov 등)일 수 있어서 정답 파일 기준으로 재면 안 된다.
    """
    out = dict.fromkeys(MANIFEST_COLS, "")
    if media != "video" or not work_path:
        return out
    m = cache.get(work_path)
    if m and m.get("status") == "ok":
        out.update(gt_width=m["width"], gt_height=m["height"], gt_resolution=resolution(m),
                   gt_fps=m.get("fps") or "", gt_duration_sec=m.get("duration_sec") or "",
                   gt_resolution_source="ffprobe 실측 (tier B)")
    s_res, s_dur_raw, _ = site_by_bucket.get(license_bucket, {}).get(str(set_id), (None, None, None))
    out["gt_resolution_site"] = s_res or ""
    sd = site_seconds(s_dur_raw)
    out["gt_duration_site_sec"] = sd if sd is not None else ""
    # 게시값은 저작물 단위라 ZIP 에 실린 변형(swf 미리보기 등)과 다를 수 있다.
    # False 인 행은 게시값으로 채점하면 정확한 추출도 오답이 된다.
    if s_res and out["gt_resolution"]:
        out["gt_resolution_site_matches_file"] = s_res == out["gt_resolution"]
    return out


def add_video_technical(df, cache: dict | None = None):
    """DataFrame(contract_work_manifest 행) 에 영상 기술 정답 열을 붙인다. 측정은 하지 않는다."""
    cache = load_cache() if cache is None else cache
    site = {b: load_site_values(b) for b in BUCKETS}
    add = [video_technical_columns(str(r["work_path"] or ""), r["media"], r["license_bucket"],
                                   str(r["set_id"]), cache, site) for _, r in df.iterrows()]
    for c in MANIFEST_COLS:
        df[c] = [a[c] for a in add]
    return df


def enrich_contract_manifest(csv_path: Path, cache: dict) -> dict:
    import pandas as pd
    _backup(csv_path, csv_path.parent)      # dataset/ 은 git 밖이라 이력이 없다
    df = pd.read_csv(csv_path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    df = add_video_technical(df, cache)
    df.to_csv(csv_path, index=False, encoding="utf-8-sig")
    v = df[df.media == "video"]
    return {"video_rows": len(v), "with_measure": int((v.gt_resolution != "").sum()),
            "with_site": int((v.gt_resolution_site != "").sum()),
            "site_matches_shipped_file": int((v.gt_resolution_site_matches_file == True).sum()),  # noqa: E712
            "site_differs_from_shipped_file": int((v.gt_resolution_site_matches_file == False).sum())}  # noqa: E712


# ---------------------------------------------------------------- main
def collect_paths(manifest_csv: Path | None) -> list:
    import openpyxl
    paths = []
    wb = openpyxl.load_workbook(KOGL_INDEX, read_only=True)
    ws = wb["originals_index"]
    it = ws.iter_rows(values_only=True)
    hdr = {n: i for i, n in enumerate(next(it))}
    for r in it:
        if r[hdr["분류"]] == "영상" and r[hdr["saved_file"]]:
            paths.append(str(KOGL_ROOT / "영상" / r[hdr["saved_file"]]))
    wb.close()
    for b in BUCKETS:
        for _, _, ps in gongu_rows(b):
            paths.extend(ps)
    if manifest_csv and manifest_csv.exists():
        with manifest_csv.open(encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                if r.get("media") == "video" and r.get("work_path"):
                    paths.append(r["work_path"])
    return paths


def main() -> int:
    ap = argparse.ArgumentParser(description="영상 기술 메타데이터 실측 → 정답 색인 기록")
    ap.add_argument("--manifest", default=str(REPO / "dataset" / "contract_work_manifest.csv"))
    ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("--probe-only", action="store_true")
    a = ap.parse_args()
    man = Path(a.manifest)

    print("[1/4] 측정", flush=True)
    cache = measure(collect_paths(man), workers=a.workers)
    if a.probe_only:
        return 0
    print("[2/4] KOGL originals_index.xlsx", flush=True)
    print("   ", update_kogl_index(cache), flush=True)
    print("[3/4] 공유마당 gongu_index.xlsx + 사이트 대조", flush=True)
    rep = update_gongu_indexes(cache)
    print(f"    대조표 {REPORT}  ({len(rep):,}행)", flush=True)
    print("[4/4] contract_work_manifest.csv", flush=True)
    if man.exists():
        print("   ", enrich_contract_manifest(man, cache), flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
