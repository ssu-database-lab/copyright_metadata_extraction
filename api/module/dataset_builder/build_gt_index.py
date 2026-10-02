"""모든 저작물의 정답을 한 행씩 모은 통합 색인 — 공유마당 + 공공누리(KOGL).

산출
----
  dataset/gt_merged_index.csv    UTF-8(BOM). 긴 설명도 잘리지 않은 원문.
  dataset/gt_merged_index.xlsx   index · summary · columns · README 시트.
                                 summary 의 집계는 전부 index 시트를 참조하는 수식이다.

행 = 저작물 1건. 첫 열 `source` 가 정답의 출처 사이트(공유마당 / 공공누리(KOGL))이고,
`gt_record_file`·`detail_url` 로 원래 정답 레코드와 사이트 페이지까지 거슬러 갈 수 있다.

입력 (읽기만 한다)
------------------
  공유마당 저작물 정답  D:/copyright_dataset_metadata/{media}/{bucket}/ground_truth.jsonl  (10개)
  공유마당 상세 URL     같은 폴더의 records.jsonl (detail_url)
  KOGL 저작물 정답      E:/kogl_originals/ground_truth.jsonl
  KOGL 상세 URL         dataset/manifest.xlsx 게시글URL
  계약서 정답           dataset/contract_work_manifest.csv

행 구성 (2026-10-02)
--------------------
  공유마당 4,834 — 전부 계약서 있음 (계약서 매니페스트 work_source gongu_gt 4,824 + gongu_fs 10)
  KOGL     1,500 — 941 건에 계약서
  KOGL        59 — 계약서만 있고 저작물 파일이 없다. KOGL download_log 에 52건 no_in_site_download,
                   7건 failed 로 남은 저작물이라 출처는 KOGL 이다. 저작물 정답 열은 비어 있다.
두 사이트의 ID 체계(공유마당 wrtSn / KOGL 원문인덱스)는 겹치지 않지만, 키는 그래도
'gongu:<id>' / 'kogl:<id>' 로 출처를 붙인다.

열 묶음
-------
  식별·추적   source, work_id, media, title, merged_key, source_site, license_bucket, detail_url, …
  work_*      저작물 정답 (계획서 11속성 + 해시태그). 값과 tier. 목록은 ' | ' 로 잇는다.
  video_*     영상 실측값(fps·재생시간)과 공유마당 게시값
  contract_*  계약서 매니페스트의 모든 열 (TTA 14속성 정답 gt_* 포함)

사용
----
  python -m api.module.dataset_builder.build_gt_index
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import re
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[3]
GONGU_ROOT = Path("/mnt/d/copyright_dataset_metadata")
KOGL_GT = Path("/mnt/e/kogl_originals/ground_truth.jsonl")
KOGL_MANIFEST = REPO / "dataset" / "manifest.xlsx"
CONTRACT_MANIFEST = REPO / "dataset" / "contract_work_manifest.csv"
OUT_CSV = REPO / "dataset" / "gt_merged_index.csv"
OUT_XLSX = REPO / "dataset" / "gt_merged_index.xlsx"

GONGU, KOGL = "공유마당", "공공누리(KOGL)"
SITE = {GONGU: "gongu.copyright.or.kr", KOGL: "www.kogl.or.kr"}
KEY = {GONGU: "gongu", KOGL: "kogl"}
MEDIA_ORDER = {"image": 0, "video": 1, "text": 2}
KOGL_DIVISION = {"image": "img", "video": "video", "text": "doc"}

WORK_ATTRS = ["제목", "저자", "설명", "라이선스 유형", "키워드", "해시태그", "해상도",
              "주요 색상", "개체 범주", "파일크기", "파일포맷", "파일 생성 날짜"]
EXCEL_CELL_MAX = 32_000          # Excel 셀 한도 32,767자 — CSV 에는 원문 그대로 남는다

IDENTITY_COLS = ["source", "work_id", "media", "title", "merged_key", "source_site",
                 "license_bucket", "detail_url", "work_file", "work_file_exists",
                 "has_work_gt", "has_contract", "gt_record_file"]
WORK_COLS = []
for _a in WORK_ATTRS:
    WORK_COLS += [f"work_{_a}", f"work_{_a}_tier"]
WORK_COLS += ["work_설명_note", "work_설명_raw", "work_키워드_batch_level", "work_키워드_batch_size",
              "work_키워드_내용태그", "work_해시태그_page", "work_해상도_source",
              "work_분류_대분류", "work_분류_장르", "work_분류_세부"]
VIDEO_COLS = ["video_fps", "video_duration_sec", "video_site_resolution",
              "video_site_duration", "video_site_aspect"]


# ---------------------------------------------------------------- 입력
def _jsonl(p: Path):
    with p.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                try:
                    yield json.loads(line)
                except json.JSONDecodeError:
                    continue


def load_gongu() -> dict:
    out = {}
    for media in ("image", "video", "text"):
        for bucket in ("expired", "donated", "ccl", "kogl"):
            p = GONGU_ROOT / media / bucket / "ground_truth.jsonl"
            if not p.exists():
                continue
            urls = {}
            rp = p.with_name("records.jsonl")
            if rp.exists():
                for r in _jsonl(rp):                       # 재개 이력 — 마지막 값이 최신
                    if r.get("detail_url"):
                        urls[str(r.get("wrtSn"))] = r["detail_url"]
            for g in _jsonl(p):
                out[str(g["id"])] = (g, str(p), urls.get(str(g["id"])))
    return out


def load_kogl() -> dict:
    import pandas as pd
    urls = {}
    if KOGL_MANIFEST.exists():
        m = pd.read_excel(KOGL_MANIFEST, dtype=str)
        urls = dict(zip(m["원문인덱스"].astype(str), m["게시글URL"]))
    return {str(g["id"]): (g, str(KOGL_GT), urls.get(str(g["id"]))) for g in _jsonl(KOGL_GT)}


def load_contracts() -> tuple[dict, list]:
    import pandas as pd
    df = pd.read_csv(CONTRACT_MANIFEST, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    cols = [c for c in df.columns if c != "set_id"]
    return {r["set_id"]: r for r in df.to_dict("records")}, cols


CONTRACT_SOURCE_NAME: dict = {}     # 색인 열 이름 → 계약서 매니페스트의 원래 열 이름


def contract_col(c: str) -> str:
    out = c if c.startswith("contract_") else f"contract_{c}"
    CONTRACT_SOURCE_NAME[out] = c
    return out


# ---------------------------------------------------------------- 행
def flat(v):
    if v is None:
        return None
    if isinstance(v, list):
        return " | ".join(str(x) for x in v) if v else None
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False) if v else None
    if isinstance(v, str) and not v.strip():
        return None
    return v


def _contract_value(v):
    if v == "":
        return None
    if v in ("True", "False"):
        return v == "True"
    return v


def work_row(source: str, g: dict, gt_file: str, url: str | None, contract: dict | None,
             ccols: list) -> dict:
    a = g.get("attributes") or {}
    ex = g.get("extra_reference") or {}
    cl = g.get("classification") or {}
    row = {
        "source": source, "work_id": str(g["id"]), "media": g.get("media"),
        "title": flat((a.get("제목") or {}).get("value")),
        "merged_key": f"{KEY[source]}:{g['id']}", "source_site": SITE[source],
        "license_bucket": g.get("license_bucket"), "detail_url": url,
        "work_file": g.get("file"), "work_file_exists": bool(g.get("file_exists")),
        "has_work_gt": True, "has_contract": contract is not None, "gt_record_file": gt_file,
    }
    for name in WORK_ATTRS:
        e = a.get(name) or {}
        row[f"work_{name}"] = flat(e.get("value"))
        row[f"work_{name}_tier"] = e.get("tier")
    d, k, h, r = a.get("설명") or {}, a.get("키워드") or {}, a.get("해시태그") or {}, a.get("해상도") or {}
    row["work_설명_note"] = d.get("note")
    row["work_설명_raw"] = flat(d.get("raw"))
    row["work_키워드_batch_level"] = bool(k.get("batch_level")) if k else None
    row["work_키워드_batch_size"] = k.get("batch_size")
    tc = k.get("tag_classes") or {}
    row["work_키워드_내용태그"] = flat([t for t, c in tc.items() if c == "내용"]) if tc else None
    row["work_해시태그_page"] = flat(h.get("page_value"))
    row["work_해상도_source"] = r.get("source")
    row["work_분류_대분류"], row["work_분류_장르"], row["work_분류_세부"] = \
        cl.get("category"), cl.get("genre"), cl.get("detail")

    is_video = g.get("media") == "video"
    site_res = None
    if is_video and source == GONGU:
        site_res = r.get("value") if r.get("tier") == "A" else ex.get("사이트_해상도")
    row.update({
        "video_fps": ex.get("실측_fps") if is_video else None,
        "video_duration_sec": ex.get("실측_재생시간_sec") if is_video else None,
        "video_site_resolution": site_res,
        "video_site_duration": ex.get("재생시간") if is_video else None,
        "video_site_aspect": ex.get("화면비율") if is_video else None,
    })
    for c in ccols:
        row[contract_col(c)] = _contract_value(contract[c]) if contract else None
    return row


def contract_only_row(sid: str, c: dict, ccols: list) -> dict:
    media = c.get("media") or c.get("media_declared") or None
    row = {k: None for k in IDENTITY_COLS + WORK_COLS + VIDEO_COLS}
    row.update({
        "source": KOGL, "work_id": sid, "media": media, "title": c.get("gt_title") or None,
        "merged_key": f"kogl:{sid}", "source_site": SITE[KOGL],
        "license_bucket": c.get("license_bucket") or None,
        "detail_url": (f"https://www.kogl.or.kr/recommend/recommendDivView.do?"
                       f"recommendIdx={sid}&division={KOGL_DIVISION.get(media, 'img')}"),
        "work_file": None, "work_file_exists": False,
        "has_work_gt": False, "has_contract": True, "gt_record_file": None,
    })
    for col in ccols:
        row[contract_col(col)] = _contract_value(c[col])
    return row


def build() -> tuple[list, list]:
    gongu, kogl = load_gongu(), load_kogl()
    contracts, ccols = load_contracts()
    columns = IDENTITY_COLS + WORK_COLS + VIDEO_COLS + [contract_col(c) for c in ccols]

    rows, used = [], set()
    for source, recs in ((GONGU, gongu), (KOGL, kogl)):
        for wid, (g, gt_file, url) in recs.items():
            c = contracts.get(wid)
            if c is not None:
                used.add(wid)
            rows.append(work_row(source, g, gt_file, url, c, ccols))
    for sid, c in contracts.items():
        if sid in used:
            continue
        if c.get("work_source"):
            raise SystemExit(f"계약서 {sid} 의 저작물({c['work_source']})이 어느 정답에도 없습니다")
        rows.append(contract_only_row(sid, c, ccols))

    keys = [r["merged_key"] for r in rows]
    if len(keys) != len(set(keys)):
        raise SystemExit("merged_key 중복")
    rows.sort(key=lambda r: (r["source"] != GONGU, MEDIA_ORDER.get(r["media"], 9),
                             str(r["license_bucket"]), r["work_id"].zfill(12)))
    return rows, columns


# ---------------------------------------------------------------- 출력
GROUP_OF = {}
for _c in IDENTITY_COLS:
    GROUP_OF[_c] = "identity"
for _c in WORK_COLS:
    GROUP_OF[_c] = "work"
for _c in VIDEO_COLS:
    GROUP_OF[_c] = "video"
GROUP_FILL = {"identity": "D9D9D9", "work": "DDEBF7", "video": "E2EFDA", "contract": "FCE4D6"}

DESCRIPTIONS = {
    "source": "Site the work's ground truth comes from: 공유마당 (gongu.copyright.or.kr) or 공공누리(KOGL) (www.kogl.or.kr).",
    "work_id": "Work ID on that site: 공유마당 wrtSn / KOGL 원문인덱스 (= contract set_id).",
    "media": "image / video / text.",
    "title": "Display title: work GT 제목, or the contract's gt_title for contract-only rows.",
    "merged_key": "Unique key: 'gongu:<id>' or 'kogl:<id>'.",
    "source_site": "Site domain.",
    "license_bucket": "Work-side bucket: 공유마당 expired / donated / ccl / kogl; KOGL works are 'kogl' (contract-only rows use the manifest bucket).",
    "detail_url": "The work's page on the source site (contract-only KOGL rows: built from the URL pattern).",
    "work_file": "Original file the work GT refers to.",
    "work_file_exists": "Whether that file exists on disk.",
    "has_work_gt": "FALSE only for the 59 KOGL works that have a contract but whose file never downloaded.",
    "has_contract": "Whether a synthetic contract exists for this work (dataset/contract_work_manifest.csv).",
    "gt_record_file": "The ground_truth.jsonl the work-side values were read from.",
    "work_설명_note": "Why a description was excluded (e.g. original-title only, empty detail page) or a remark on it.",
    "work_설명_raw": "Collected description text that was excluded from scoring (kept for review).",
    "work_키워드_batch_level": "TRUE when 5+ works share exactly this tag list — collection tags, not content of this work.",
    "work_키워드_batch_size": "Number of works sharing the tag list.",
    "work_키워드_내용태그": "KOGL only: the tags classified as content (not agency / format / collection / title word).",
    "work_해시태그_page": "KOGL only: the tags the detail page actually shows.",
    "work_해상도_source": "Where 해상도 came from: site-published value (tier A) or measurement (PIL / ffprobe, tier B).",
    "work_분류_대분류": "공유마당 classification, top level (영상 / 미술 / 사진 / 어문).",
    "work_분류_장르": "공유마당 classification, genre.",
    "work_분류_세부": "공유마당 classification, detail.",
    "video_fps": "ffprobe-measured frame rate (avg_frame_rate).",
    "video_duration_sec": "ffprobe-measured duration in seconds (SWF: packets ÷ fps).",
    "video_site_resolution": "공유마당 only: the resolution the site publishes (may describe another file of the same work).",
    "video_site_duration": "공유마당 only: the duration the site publishes (floored to whole seconds).",
    "video_site_aspect": "공유마당 only: the aspect ratio the site publishes.",
}


def describe(col: str) -> tuple[str, str]:
    if col in DESCRIPTIONS:
        return GROUP_OF.get(col, "contract"), DESCRIPTIONS[col]
    if col.startswith("work_") and col.endswith("_tier"):
        return "work", (f"Tier of work_{col[5:-5]}: A external · B measured from the file · "
                        "C no answer · D different meaning (skipped) · N/A · 제외 excluded.")
    if col.startswith("work_"):
        return "work", (f"Work-side ground truth '{col[5:]}' (lists joined with ' | '). "
                        + ("Reference only — not a scored attribute." if col == "work_해시태그" else ""))
    if col.startswith("contract_gt_"):
        return "contract", (f"Contract-side ground truth (contract manifest column "
                            f"'{CONTRACT_SOURCE_NAME.get(col, col)}').")
    return "contract", f"Contract manifest column '{CONTRACT_SOURCE_NAME.get(col, col)}'."


def write_csv(rows, columns):
    import pandas as pd
    pd.DataFrame(rows, columns=columns).to_csv(OUT_CSV, index=False, encoding="utf-8-sig")


def write_xlsx(rows, columns) -> dict:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.indexed_list import IndexedList

    wb = Workbook()
    wb._fonts = IndexedList([Font(name="Arial", sz=10)])          # 빈 셀까지 Arial
    wb._named_styles["Normal"].font = Font(name="Arial", sz=10)
    # summary 는 수식뿐이고 openpyxl 은 계산값을 저장하지 못한다 — 열 때 Excel 이 전부 계산하게 한다
    wb.calculation.fullCalcOnLoad = True
    bold = Font(name="Arial", sz=10, bold=True)
    title_font = Font(name="Arial", sz=13, bold=True)
    stats = {"truncated": 0, "illegal_chars": 0}

    # ---- index
    ws = wb.active
    ws.title = "index"
    ws.append(columns)
    for j, c in enumerate(columns, 1):
        cell = ws.cell(1, j)
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor=GROUP_FILL[GROUP_OF.get(c, "contract")])
    for i, r in enumerate(rows, 2):
        for j, c in enumerate(columns, 1):
            v = r.get(c)
            if v is None:
                continue
            if isinstance(v, str):
                if ILLEGAL_CHARACTERS_RE.search(v):
                    v = ILLEGAL_CHARACTERS_RE.sub("", v)
                    stats["illegal_chars"] += 1
                if len(v) > EXCEL_CELL_MAX:
                    v = v[:EXCEL_CELL_MAX] + f" …[truncated; full text ({len(r[c]):,} chars) in the CSV]"
                    stats["truncated"] += 1
            cell = ws.cell(i, j, v)
            if isinstance(v, str) and v.startswith("="):
                cell.data_type = "s"                  # 정답 문자열이 수식으로 해석되지 않게
    for j, c in enumerate(columns, 1):
        sample = [len(str(r.get(c) or "")) for r in rows[:400]]
        width = max(len(c) + 2, min(45, max(sample + [8]) + 2))
        ws.column_dimensions[get_column_letter(j)].width = 42 if c == "title" else width
    ws.freeze_panes = "E2"                            # source · work_id · media · title 고정
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    col_letter = {c: get_column_letter(j) for j, c in enumerate(columns, 1)}

    def rng(c):
        return f"index!${col_letter[c]}:${col_letter[c]}"

    # ---- summary (전부 index 를 참조하는 수식)
    sm = wb.create_sheet("summary")
    sm["A1"] = "Merged ground-truth index — summary (all counts are formulas over the 'index' sheet)"
    sm["A1"].font = title_font
    media = ["image", "video", "text"]
    sources = [GONGU, KOGL]

    def table(top, title, extra=None):
        sm.cell(top, 1, title).font = bold
        sm.cell(top + 1, 1, "source").font = bold
        for k, m in enumerate(media, 2):
            sm.cell(top + 1, k, m).font = bold
        sm.cell(top + 1, 5, "total").font = bold
        for s_i, s in enumerate(sources):
            rr = top + 2 + s_i
            sm.cell(rr, 1, s)
            for k, m in enumerate(media, 2):
                crit = f'{rng("source")},$A{rr},{rng("media")},{get_column_letter(k)}${top + 1}'
                if extra:
                    crit += f",{extra}"
                sm.cell(rr, k, f"=COUNTIFS({crit})")
            sm.cell(rr, 5, f"=SUM(B{rr}:D{rr})")
        tr = top + 2 + len(sources)
        sm.cell(tr, 1, "total").font = bold
        for k in range(2, 6):
            L = get_column_letter(k)
            sm.cell(tr, k, f"=SUM({L}{top + 2}:{L}{tr - 1})").font = bold
        return tr + 2

    nxt = table(3, "Works by source × media")
    nxt = table(nxt, "…of which have a contract", f'{rng("has_contract")},TRUE')
    nxt = table(nxt, "…of which have work-side ground truth", f'{rng("has_work_gt")},TRUE')
    nxt = table(nxt, "…of which have a scorable description (work_설명 filled)", f'{rng("work_설명")},"<>"')
    nxt = table(nxt, "…of which have collection-level keywords (batch_level TRUE)",
                f'{rng("work_키워드_batch_level")},TRUE')

    sm.cell(nxt, 1, "Work-side GT coverage per attribute (non-empty value)").font = bold
    sm.cell(nxt + 1, 1, "attribute").font = bold
    for k, s in enumerate(sources, 2):
        sm.cell(nxt + 1, k, s).font = bold
    sm.cell(nxt + 1, 4, "of which tier A").font = bold
    for i_a, a in enumerate(WORK_ATTRS):
        rr = nxt + 2 + i_a
        sm.cell(rr, 1, a)
        for k, s in enumerate(sources, 2):
            sm.cell(rr, k, f'=COUNTIFS({rng("source")},{get_column_letter(k)}${nxt + 1},{rng("work_" + a)},"<>")')
        sm.cell(rr, 4, f'=COUNTIFS({rng("work_" + a)},"<>",{rng("work_" + a + "_tier")},"A")')
    sm.column_dimensions["A"].width = 62
    for L in "BCDE":
        sm.column_dimensions[L].width = 18

    # ---- columns (data dictionary)
    cd = wb.create_sheet("columns")
    cd.append(["column", "group", "description"])
    for j in range(1, 4):
        cd.cell(1, j).font = bold
    for c in columns:
        grp, desc = describe(c)
        cd.append([c, grp, desc])
        cd.cell(cd.max_row, 1).fill = PatternFill("solid", fgColor=GROUP_FILL[grp])
    cd.column_dimensions["A"].width = 38
    cd.column_dimensions["B"].width = 11
    cd.column_dimensions["C"].width = 120
    cd.freeze_panes = "A2"

    # ---- README
    rd = wb.create_sheet("README")
    lines = [
        ("Merged ground-truth index — 공유마당 + 공공누리(KOGL)", title_font),
        (f"Built {_dt.date.today().isoformat()} by api/module/dataset_builder/build_gt_index.py", None),
        ("", None),
        ("One row per work. The first column, 'source', says which site the ground truth comes from;", None),
        ("'gt_record_file' and 'detail_url' trace each row back to its GT record and the site page.", None),
        ("Counts are on the 'summary' sheet; every column is explained on the 'columns' sheet.", None),
        ("", None),
        ("Column groups (header colours)", bold),
        ("  grey   — identity and traceability", None),
        ("  blue   — work_*: work-side GT (the 11 plan attributes + 해시태그), value and tier", None),
        ("  green  — video_*: ffprobe measurements and 공유마당 published values", None),
        ("  orange — contract_*: every column of dataset/contract_work_manifest.csv, incl. the TTA-14 gt_* answers", None),
        ("", None),
        ("Sources read (not modified)", bold),
        (f"  {GONGU_ROOT}/<media>/<bucket>/ground_truth.jsonl (+ records.jsonl for detail_url)", None),
        (f"  {KOGL_GT}", None),
        (f"  {KOGL_MANIFEST.relative_to(REPO)} (KOGL 게시글URL)", None),
        (f"  {CONTRACT_MANIFEST.relative_to(REPO)}", None),
        ("", None),
        ("Read before using the scores", bold),
        ("  • Tier: A = external answer · B = measured from the file (checks the value survives the pipeline,"
         " not accuracy) · C = no answer · D = different meaning (skipped) · 제외 = excluded (value empty, text in work_설명_raw).", None),
        ("  • Rows with has_work_gt = FALSE are KOGL works that have a contract but whose file never downloaded"
         " (52 no download on the site, 7 failed) — work_* columns are empty.", None),
        ("  • work_키워드_batch_level = TRUE marks tag lists shared by 5+ works: collection tags, not content."
         " Report keyword results for these separately.", None),
        ("  • 해시태그 is reference only: neither the plan-11 nor the TTA-14 list scores it.", None),
        ("  • 99 KOGL image detail pages return an identical empty page; their description is 'unknown', not absent.", None),
        (f"  • Excel cells hold at most 32,767 characters: {stats['truncated']} long texts are truncated here"
         " (marked '…[truncated…]'); the CSV next to this file has them in full.", None),
        ("", None),
        ("Rebuild", bold),
        ("  python -m api.module.dataset_builder.build_gt_index", None),
    ]
    for text, font in lines:
        rd.append([text])
        if font:
            rd.cell(rd.max_row, 1).font = font
    rd.column_dimensions["A"].width = 150
    for row in rd.iter_rows():
        row[0].alignment = Alignment(wrap_text=False)

    wb.save(OUT_XLSX)
    return stats


def main() -> int:
    argparse.ArgumentParser(description="공유마당 + KOGL 통합 정답 색인").parse_args()
    rows, columns = build()
    write_csv(rows, columns)
    stats = write_xlsx(rows, columns)
    from collections import Counter
    c = Counter((r["source"], r["media"]) for r in rows)
    print(f"{len(rows):,} rows × {len(columns)} columns")
    for k, v in sorted(c.items()):
        print(f"  {k[0]:<14} {k[1]:<6} {v:>5}")
    print(f"  has_contract {sum(r['has_contract'] for r in rows):,} · "
          f"no work GT {sum(not r['has_work_gt'] for r in rows)}")
    print(f"  excel: truncated {stats['truncated']} · illegal-char cells cleaned {stats['illegal_chars']}")
    print(f"→ {OUT_CSV}\n→ {OUT_XLSX}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
